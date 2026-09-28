"""Word (.docx) から本文・Legend・Table を構造化抽出."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document


_LEGEND_RE = re.compile(
    r"^(Figure|Fig\.?|Supplementary Figure|Table)\s*S?\d+",
    re.IGNORECASE,
)
_N_RE = re.compile(r"\bn\s*=\s*(\d+)", re.IGNORECASE)


@dataclass
class LegendBlock:
    label: str
    text: str
    sample_sizes: list[int] = field(default_factory=list)


@dataclass
class ManuscriptSections:
    paragraphs: list[str]
    legends: list[LegendBlock]
    tables: list[list[list[str]]]
    raw_text: str

    def sample_size_mentions(self) -> list[tuple[str, int]]:
        """(コンテキスト, n) のリスト."""
        out: list[tuple[str, int]] = []
        for leg in self.legends:
            for n in leg.sample_sizes:
                out.append((leg.label, n))
        for m in _N_RE.finditer(self.raw_text):
            start = max(0, m.start() - 40)
            end = min(len(self.raw_text), m.end() + 40)
            out.append((self.raw_text[start:end].replace("\n", " "), int(m.group(1))))
        return out


def extract_docx(path: Path | str) -> ManuscriptSections:
    doc = Document(str(path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    legends: list[LegendBlock] = []
    for text in paragraphs:
        if _LEGEND_RE.match(text):
            label = text.split(".", 1)[0].strip() if "." in text[:40] else text[:40]
            ns = [int(x) for x in _N_RE.findall(text)]
            legends.append(LegendBlock(label=label, text=text, sample_sizes=ns))

    tables: list[list[list[str]]] = []
    for table in doc.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        tables.append(rows)

    raw_text = "\n".join(paragraphs)
    return ManuscriptSections(
        paragraphs=paragraphs,
        legends=legends,
        tables=tables,
        raw_text=raw_text,
    )
