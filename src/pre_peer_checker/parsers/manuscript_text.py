"""原稿テキストを段落リストにする（Word / PDF 共通の入口）.

Legend・本文・References の各パーサは段落リストを受け取るので、PDF でも
Word と同じ形に揃えれば同じチェックが動く。PDF は埋め込みテキストを読み、
テキスト層の無いページだけ OCR（PyMuPDF + Tesseract があれば）で補う。
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

_HEADING_RE = re.compile(
    r"^(?:"
    r"abstract|summary|introduction|results?|discussion|conclusions?|"
    r"(?:materials?\s+and\s+)?methods|online\s+methods|star\W*methods|"
    r"method\s+details|experimental\s+(?:procedures|model\s+and\s+(?:study\s+)?subject\s+details)|"
    r"key\s+resources\s+table|resource\s+availability|"
    r"quantification\s+and\s+statistical\s+analysis|statistical\s+analysis|"
    r"statistics(?:\s+and\s+reproducibility)?|reporting\s+summary|"
    r"references?|bibliography|references\s+and\s+notes|"
    r"acknowledg?e?ments?|author\s+contributions?|funding|"
    r"competing\s+interests?|declaration\s+of\s+interests|"
    r"data\s+availability|code\s+availability|additional\s+information|"
    r"supplementa(?:l|ry)\s+(?:information|materials?|note)|"
    r"figure\s+legends?|table\s+legends?|highlights|graphical\s+abstract"
    r")\s*$",
    re.IGNORECASE,
)
_REF_HEAD_RE = re.compile(r"^(?:references?|bibliography|references\s+and\s+notes)\s*$", re.I)
_LEGEND_HEAD_RE = re.compile(
    r"^(?:Figure|Fig\.?|Supplementary\s+Figure|Extended\s+Data\s+Fig\.?)\s*S?\d+\s*[.:|]"
)
_PANEL_OPEN_RE = re.compile(r"^\([A-Z]\d?(?:\s*(?:[-–,]|and)\s*[A-Z]\d?)*\)")
_BRACKET_CITE_RE = re.compile(r"\[\d{1,3}(?:\s*[,–\-]\s*\d{1,3})*\]")
_SENT_END_RE = re.compile(r"[.!?:;)\]]$")
_LEGEND_CONT_RE = re.compile(r"^\(?(?:figure\s+)?legend\s+continued.*\)?$", re.I)
_PAGE_NUM_RE = re.compile(r"^(?:page\s+)?\d{1,4}(?:\s*(?:of|/)\s*\d{1,4})?$", re.I)
_REF_ENTRY_START_RE = re.compile(r"^(?:\[\d{1,3}\]|\d{1,3}\.)\s+\S")
_SUP_CITE_RE = re.compile(r"^\d{1,3}(?:\s*[,–\-]\s*\d{1,3})*$")
# Superscripts right after these are units/exponents (µm2, 10^6), not citations.
_UNIT_TAIL_RE = re.compile(
    r"(?:\d|(?:^|[^A-Za-z])(?:[kcmdµμn]?m|[kmµμn]?g|[mµμ]?[lL]|s|h|min|°C))$"
)

_OCR_DPI = 300
_MIN_TEXT_CHARS = 40


@dataclass
class _Block:
    page: int
    text: str
    size: float
    n_lines: int
    n_words: int


@dataclass
class PdfText:
    """PDF から復元した段落と、読み取り状況."""

    paragraphs: list[str] = field(default_factory=list)
    n_pages: int = 0
    ocr_pages: list[int] = field(default_factory=list)
    image_only_pages: list[int] = field(default_factory=list)
    legends_relocated: int = 0


def _span_text(spans: list[dict], line_size: float) -> str:
    """Join spans; superscript citation numbers become ``[1–4]`` like Word manuscripts."""
    out: list[str] = []
    sup: list[str] = []

    def _flush_sup() -> None:
        if not sup:
            return
        blob = "".join(sup).strip()
        prev = "".join(out).rstrip()
        if (
            _SUP_CITE_RE.match(blob)
            and prev
            and not _UNIT_TAIL_RE.search(prev)
            and re.search(r"[A-Za-z.,;:)\]]$", prev)
        ):
            out.append(f"[{blob}]")
        else:
            out.append(blob)
        sup.clear()

    for s in spans:
        t = s.get("text") or ""
        is_sup = bool(int(s.get("flags") or 0) & 1) and float(s.get("size") or 0) < line_size * 0.85
        if is_sup:
            sup.append(t)
            continue
        _flush_sup()
        out.append(t)
    _flush_sup()
    return "".join(out)


def _join_lines(lines: list[str]) -> str:
    text = ""
    for ln in lines:
        ln = re.sub(r"\s+", " ", ln).strip()
        if not ln:
            continue
        if not text:
            text = ln
        elif text.endswith("-") and ln[:1].islower() and not text.endswith(" -"):
            text = text[:-1] + ln
        else:
            text = f"{text} {ln}"
    return text


def _page_blocks(page_dict: dict, page_no: int) -> list[_Block]:
    blocks: list[_Block] = []
    for b in page_dict.get("blocks") or []:
        if b.get("type") != 0:
            continue
        sizes: Counter[float] = Counter()
        texts: list[str] = []
        for ln in b.get("lines") or []:
            spans = ln.get("spans") or []
            if not spans:
                continue
            line_size = max(float(s.get("size") or 0) for s in spans)
            for s in spans:
                sizes[round(float(s.get("size") or 0), 1)] += len((s.get("text") or "").strip())
            texts.append(_span_text(spans, line_size))
        texts = [t for t in texts if t.strip()]
        if not texts:
            continue
        size = sizes.most_common(1)[0][0] if sizes else 0.0
        text = _join_lines(texts)
        blocks.append(
            _Block(
                page=page_no,
                text=text,
                size=size,
                n_lines=len(texts),
                n_words=len(text.split()),
            )
        )
    return blocks


@lru_cache(maxsize=1)
def ocr_available() -> bool:
    """True when PyMuPDF can find a Tesseract installation (language data)."""
    import pymupdf

    try:
        return bool(pymupdf.get_tessdata())
    except Exception:  # noqa: BLE001
        return False


def _ocr_page_dict(page, flags: int) -> dict:
    tp = page.get_textpage_ocr(flags=flags, dpi=_OCR_DPI, full=True)
    return page.get_text("dict", flags=flags, textpage=tp)


def _is_heading(text: str) -> bool:
    return bool(_HEADING_RE.match(text.strip()))


def _split_heading(block: _Block) -> list[_Block]:
    """Detach a section heading merged into the first line of a body block."""
    for head_words in (1, 2, 3, 4, 5, 6):
        parts = block.text.split(" ", head_words)
        if len(parts) <= head_words:
            break
        head = " ".join(parts[:head_words])
        nxt = parts[head_words][:1]
        if _is_heading(head) and (nxt.isupper() or nxt.isdigit() or nxt == "["):
            rest = parts[head_words]
            return [
                _Block(block.page, head, block.size, 1, head_words),
                _Block(block.page, rest, block.size, max(1, block.n_lines - 1), len(rest.split())),
            ]
    return [block]


def _looks_like_prose(block: _Block) -> bool:
    """Drop figure-internal labels (axis ticks, panel letters, stacked short words)."""
    lines = max(block.n_lines, 1)
    if block.n_words >= 8 and (block.n_words / lines >= 4 or len(block.text) / lines >= 22):
        return True
    if block.n_words >= 6 and _REF_ENTRY_START_RE.match(block.text):
        return True
    return block.n_lines == 1 and block.n_words >= 5 and block.text.rstrip().endswith((".", ":"))


def _running_keys(pages: list[list[_Block]]) -> set[str]:
    """Header/footer lines repeated across pages (journal name, DOI, page numbers)."""
    n_pages = len(pages)
    if n_pages < 3:
        return set()
    counts: Counter[str] = Counter()
    for blocks in pages:
        counts.update({_norm_key(b.text) for b in blocks if b.n_words <= 30})
    threshold = max(3, int(n_pages * 0.2))
    return {k for k, c in counts.items() if c >= threshold}


def _norm_key(text: str) -> str:
    """Digits → ``#`` and page numbers at either end dropped, so odd/even footers match."""
    key = re.sub(r"\d+", "#", re.sub(r"\s+", " ", text.strip().lower()))
    return re.sub(r"^(?:[a-z]?#\s+)+|(?:\s+[a-z]?#)+$", "", key)


_NAME_MARK_RE = re.compile(
    r"(?:[A-Z][\w’'\-]+\.?\s+){1,3}[A-Z][\w’'\-]+,?\s?\[\d{1,3}(?:\s*[,–\-]\s*\d{1,3})*\]"
)


def _strip_affiliation_marks(text: str) -> str:
    """Author lists carry superscript affiliation numbers, not citations."""
    cites = _BRACKET_CITE_RE.findall(text)
    if len(cites) < 3:
        return text
    name_marks = len(_NAME_MARK_RE.findall(text))
    if name_marks >= 3 and name_marks * 2 >= len(cites):
        return _BRACKET_CITE_RE.sub("", text)
    return text


def _split_reference_entries(text: str) -> list[str]:
    """Split a block holding several numbered references into one entry per paragraph."""
    parts = re.split(r"\s(?=(?:\[\d{1,3}\]|\d{1,3}\.)\s+[A-Z])", text)
    return [p.strip() for p in parts if p.strip()]


def _assemble(pages: list[list[_Block]]) -> tuple[list[str], int]:
    running = _running_keys(pages)
    kept: list[_Block] = []
    for blocks in pages:
        for b in blocks:
            t = b.text.strip()
            if not t or _PAGE_NUM_RE.match(t) or _LEGEND_CONT_RE.match(t):
                continue
            if _norm_key(t) in running:
                continue
            for part in _split_heading(b):
                if _is_heading(part.text) or _LEGEND_HEAD_RE.match(part.text) or _looks_like_prose(part):
                    kept.append(part)

    size_chars: Counter[float] = Counter()
    for b in kept:
        size_chars[round(b.size * 2) / 2] += len(b.text)
    body_size = size_chars.most_common(1)[0][0] if size_chars else 0.0

    def _smaller(b: _Block) -> bool:
        return body_size > 0 and b.size <= body_size - 0.4

    # Typeset journals float legends (in a smaller font) between body paragraphs.
    # Move them into a "Figure legends" section like a Word manuscript so body
    # paragraphs stay contiguous and each legend ends at the next legend.
    legends: list[list[_Block]] = []
    body: list[_Block] = []
    open_leg: list[_Block] | None = None
    last_was_legend = False
    for b in kept:
        is_head = bool(_LEGEND_HEAD_RE.match(b.text)) and _smaller(b)
        if is_head:
            open_leg = [b]
            legends.append(open_leg)
            last_was_legend = True
            continue
        if open_leg is not None and abs(b.size - open_leg[0].size) <= 0.3 and not _is_heading(b.text):
            last = open_leg[-1]
            follows = last_was_legend and b.page == last.page
            panel_cont = b.page in (last.page, last.page + 1) and _PANEL_OPEN_RE.match(b.text)
            if follows or panel_cont:
                open_leg.append(b)
                last_was_legend = True
                continue
        last_was_legend = False
        body.append(b)

    # Rejoin paragraphs cut by a column or page break (no sentence end, next starts lowercase).
    merged: list[_Block] = []
    for b in body:
        prev = merged[-1] if merged else None
        if (
            prev is not None
            and not _is_heading(prev.text)
            and not _is_heading(b.text)
            and abs(prev.size - b.size) <= 0.3
            and not _SENT_END_RE.search(prev.text)
            and b.text[:1].islower()
        ):
            merged[-1] = _Block(
                prev.page,
                _join_lines([prev.text, b.text]),
                prev.size,
                prev.n_lines + b.n_lines,
                prev.n_words + b.n_words,
            )
            continue
        merged.append(b)

    legend_paras = [_join_lines([x.text for x in leg]) for leg in legends]
    paragraphs: list[str] = []
    in_refs = False
    legends_placed = False
    for b in merged:
        text = _strip_affiliation_marks(b.text)
        if _is_heading(text):
            is_ref = bool(_REF_HEAD_RE.match(text))
            if is_ref and legend_paras and not legends_placed:
                paragraphs.append("Figure legends")
                paragraphs.extend(legend_paras)
                legends_placed = True
            in_refs = is_ref
            paragraphs.append(text)
            continue
        if in_refs and _REF_ENTRY_START_RE.match(text):
            paragraphs.extend(_split_reference_entries(text))
        else:
            paragraphs.append(text)
    if legend_paras and not legends_placed:
        paragraphs.append("Figure legends")
        paragraphs.extend(legend_paras)
    return paragraphs, len(legends)


def extract_pdf_text(path: Path | str, *, ocr: bool = True) -> PdfText:
    """Rebuild manuscript paragraphs from a PDF (embedded text; OCR for image-only pages)."""
    import pymupdf

    flags = pymupdf.TEXTFLAGS_DICT & ~pymupdf.TEXT_PRESERVE_LIGATURES & ~pymupdf.TEXT_PRESERVE_IMAGES
    result = PdfText()
    pages: list[list[_Block]] = []
    ocr_ok = ocr and ocr_available()
    with pymupdf.open(str(path)) as doc:
        result.n_pages = doc.page_count
        for i, page in enumerate(doc):
            page_dict = page.get_text("dict", flags=flags)
            chars = sum(
                len(s.get("text") or "")
                for b in page_dict.get("blocks") or []
                for ln in b.get("lines") or []
                for s in ln.get("spans") or []
            )
            if chars < _MIN_TEXT_CHARS and page.get_images(full=False):
                if ocr_ok:
                    try:
                        page_dict = _ocr_page_dict(page, flags)
                        result.ocr_pages.append(i + 1)
                    except Exception:  # noqa: BLE001 — Tesseract not installed
                        ocr_ok = False
                        result.image_only_pages.append(i + 1)
                else:
                    result.image_only_pages.append(i + 1)
            pages.append(_page_blocks(page_dict, i + 1))
    result.paragraphs, result.legends_relocated = _assemble(pages)
    return result


@lru_cache(maxsize=16)
def _pdf_text_cached(path: str, mtime: float, size: int) -> PdfText:
    return extract_pdf_text(path)


def pdf_text(path: Path | str) -> PdfText:
    p = Path(path)
    st = p.stat()
    return _pdf_text_cached(str(p.resolve()), st.st_mtime, st.st_size)


def docx_paragraphs(path: Path | str) -> list[str]:
    from docx import Document

    doc = Document(str(path))
    return [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]


def manuscript_paragraphs(path: Path | str) -> list[str]:
    """Ordered non-empty paragraphs of a Word or PDF manuscript."""
    if Path(path).suffix.lower() == ".pdf":
        return list(pdf_text(path).paragraphs)
    return docx_paragraphs(path)


def is_manuscript_pdf(path: Path | str, *, min_chars: int = 3000) -> bool:
    """True when the PDF carries manuscript prose (not a figure-only or plot PDF)."""
    p = Path(path)
    name = p.name.lower()
    if name.startswith(("rplot", "graph")):
        return False
    try:
        pt = pdf_text(p)
    except Exception:  # noqa: BLE001
        return False
    if pt.n_pages < 2:
        return False
    paras = pt.paragraphs
    signals = sum(1 for t in paras if _is_heading(t) or _LEGEND_HEAD_RE.match(t))
    return sum(len(t) for t in paras) >= min_chars and signals >= 2


def manuscript_source_artifact(kind: str, paths: list[Path]) -> dict:
    """Where the manuscript text came from, for coverage / report notes."""
    art: dict = {"kind": kind, "paths": [str(p) for p in paths]}
    if kind == "pdf":
        ocr: list[dict] = []
        image_only: list[dict] = []
        for p in paths:
            pt = pdf_text(p)
            if pt.ocr_pages:
                ocr.append({"path": str(p), "pages": pt.ocr_pages})
            if pt.image_only_pages:
                image_only.append({"path": str(p), "pages": pt.image_only_pages})
        art["ocr_pages"] = ocr
        art["image_only_pages"] = image_only
    return art


def select_manuscript_pdfs(
    pdfs: list[Path],
    *,
    exclude_under: list[Path] | None = None,
) -> list[Path]:
    """Pick the PDF that reads as the manuscript (most prose) when no Word file is given."""
    excluded = [Path(x).resolve() for x in (exclude_under or [])]

    def _excluded(p: Path) -> bool:
        rp = p.resolve()
        return any(rp == e or e in rp.parents for e in excluded)

    candidates = [p for p in pdfs if not _excluded(p) and is_manuscript_pdf(p)]
    if not candidates:
        return []
    best = max(candidates, key=lambda p: sum(len(t) for t in pdf_text(p).paragraphs))
    return [best]
