"""原稿内のパネル単位の画像使い回し（Word 原稿・PDF 原稿の双方）."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from panel_fixtures import page as _page
from panel_fixtures import photo as _photo
from PIL import Image

from pre_peer_checker.imaging.panel_reuse import scan_internal_panel_reuse
from pre_peer_checker.warnings import WarningTag

pytestmark = pytest.mark.usefixtures()


def _reuse_warnings(result) -> list:
    return [
        w
        for w in result.warnings
        if w.tag == WarningTag.IMAGE_REUSE and w.metadata.get("internal_reuse")
    ]


def _degraded(img: Image.Image, side: int = 240, quality: int = 70) -> Image.Image:
    small = img.resize((side, side), Image.Resampling.BILINEAR)
    buf = io.BytesIO()
    small.save(buf, format="JPEG", quality=quality)
    return Image.open(io.BytesIO(buf.getvalue())).convert("RGB")


def test_same_photo_in_two_figures_is_flagged(tmp_path: Path):
    pytest.importorskip("lightglue")
    shared = _photo(21)
    _page([_photo(22), shared]).save(tmp_path / "fig2.png")
    _page([_photo(23), _degraded(shared)]).save(tmp_path / "fig5.png")

    result = scan_internal_panel_reuse(
        [tmp_path / "fig2.png", tmp_path / "fig5.png"], prefer_dino=False
    )
    warns = _reuse_warnings(result)
    assert len(warns) == 1
    assert warns[0].metadata["same_file"] is False
    assert {Path(s).name for s in warns[0].sources} == {"fig2.png", "fig5.png"}


def test_same_photo_twice_in_one_figure_is_flagged(tmp_path: Path):
    pytest.importorskip("lightglue")
    shared = _photo(24)
    _page([shared, _photo(25), _degraded(shared)]).save(tmp_path / "fig3.png")

    result = scan_internal_panel_reuse([tmp_path / "fig3.png"], prefer_dino=False)
    warns = _reuse_warnings(result)
    assert len(warns) == 1
    assert warns[0].metadata["same_file"] is True


def test_distinct_photos_are_not_flagged(tmp_path: Path):
    pytest.importorskip("lightglue")
    _page([_photo(31), _photo(32), _photo(33)]).save(tmp_path / "fig1.png")
    _page([_photo(34), _photo(35)]).save(tmp_path / "fig4.png")

    result = scan_internal_panel_reuse(
        [tmp_path / "fig1.png", tmp_path / "fig4.png"], prefer_dino=False
    )
    assert _reuse_warnings(result) == []


def test_same_figure_in_two_formats_is_not_a_reuse_warning(tmp_path: Path):
    """PDF 埋め込みと書き出し PNG のように同じ図が 2 ファイルある場合は警告にしない."""
    pytest.importorskip("lightglue")
    composite = _page([_photo(41), _photo(42), _photo(43)])
    composite.save(tmp_path / "fig6.png")
    composite.save(tmp_path / "fig6_export.jpg", quality=88)

    result = scan_internal_panel_reuse(
        [tmp_path / "fig6.png", tmp_path / "fig6_export.jpg"], prefer_dino=False
    )
    assert _reuse_warnings(result) == []
    assert len(result.artifacts.get("duplicate_file_pairs") or []) == 1


@pytest.mark.parametrize("manuscript", ["word", "pdf"])
def test_run_verification_detects_reuse_for_word_and_pdf(tmp_path: Path, manuscript: str):
    pytest.importorskip("lightglue")
    from pre_peer_checker.pipeline.orchestrator import run_verification

    shared = _photo(51)
    composite = _page([shared, _photo(52), _degraded(shared)])
    ms = tmp_path / manuscript
    ms.mkdir()

    if manuscript == "word":
        from docx import Document

        buf = io.BytesIO()
        composite.save(buf, format="PNG")
        doc = Document()
        doc.add_paragraph("Figure 3. Eye discs across conditions. n=12 (F).")
        doc.add_picture(io.BytesIO(buf.getvalue()))
        doc.save(ms / "SuppInfo.docx")
    else:
        composite.save(ms / "FigS1.pdf", "PDF", resolution=150)

    result = run_verification([ms])
    warns = _reuse_warnings(result)
    assert len(warns) == 1
    assert warns[0].metadata["same_file"] is True
    assert result.artifacts["internal_panel_reuse"]["reuse_source_pairs"] == 1
