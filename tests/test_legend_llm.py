"""Legend LLM backend wiring (rules main; MLX/CUDA optional)."""

from __future__ import annotations

import json
from pathlib import Path

from pre_peer_checker.llm.backend import (
    CallableBackend,
    MLXBackend,
    TransformersBackend,
    probe_backends,
    select_backend,
)
from pre_peer_checker.llm.legend_extract import (
    extract_legend_json_hybrid,
    extract_legends_with_backend,
)
from pre_peer_checker.llm.verifier import LocalManuscriptVerifier
from pre_peer_checker.pipeline.orchestrator import run_verification


def test_probe_backends_returns_info():
    infos = probe_backends()
    assert len(infos) == 2
    names = {i.name for i in infos}
    assert "mlx" in names and "transformers" in names
    assert all(isinstance(i.available, bool) for i in infos)


def test_select_backend_none():
    assert select_backend("none") is None


def test_hybrid_with_callable_backend():
    def fake(prompt: str) -> str:
        return json.dumps(
            {
                "figure": "Figure 1",
                "panels": [{"panel": "F", "n": 12, "groups": ["WT"], "notes": ""}],
                "tests": ["Welch's t-test"],
                "p_values": [0.01],
                "citation": {
                    "mentioned": True,
                    "reproduced_from": "Smith 2019",
                    "spans": ["reproduced from"],
                },
                "raw_excerpt": "x",
                "extractor": "llm",
            }
        )

    text = "Figure 1. n=11 (F). No citation here."
    # rules alone: n=11, no citation
    rules = extract_legend_json_hybrid(text, figure_hint="Figure 1")
    assert rules.extractor == "rules"
    assert any(p.n == 11 for p in rules.panels)

    hybrid = extract_legend_json_hybrid(
        text, figure_hint="Figure 1", llm_generate=fake, prefer_llm=True
    )
    assert hybrid.extractor in {"llm", "llm+rules"}
    assert any(p.n == 12 for p in hybrid.panels)
    assert hybrid.citation.mentioned
    # Rules lock: explicit rule n=11 (F) survives; LLM may add another row.
    assert any(p.n == 11 and p.panel == "F" and not p.groups for p in hybrid.panels)


def test_verifier_with_stub_backend():
    be = CallableBackend(lambda p: "ok-verify", name="stub")
    v = LocalManuscriptVerifier(backend=be)
    assert v.verify_consistency("legend", {"n": 1}) == "ok-verify"
    assert v.backend_info() is not None
    assert v.backend_info().name == "stub"


def test_extract_legends_with_backend_disabled(tmp_path: Path):
    from docx import Document

    docx = tmp_path / "m.docx"
    doc = Document()
    doc.add_paragraph("Figure 1. n=6 (A). Welch's t-test.")
    doc.save(docx)
    legs, meta = extract_legends_with_backend(docx, enabled=False)
    assert legs and legs[0].extractor == "rules"
    assert meta["legend_llm_enabled"] is False


def test_orchestrator_legend_llm_flag_without_backend(tmp_path: Path):
    """When --legend-llm is on but no backend installed, rules still work."""
    from docx import Document

    ms = tmp_path / "ms"
    ms.mkdir()
    doc = Document()
    doc.add_paragraph("Figure 1. Clone size. n=5 (A1).")
    doc.save(ms / "main.docx")
    result = run_verification([ms], legend_llm=True, legend_llm_prefer="none")
    assert result.artifacts.get("legend_json")
    meta0 = result.artifacts["legend_llm"][0]
    assert meta0.get("legend_llm_enabled") is True
    assert meta0.get("backend") == "none" or meta0.get("note")

def test_mlx_transformers_available_are_bool():
    assert isinstance(MLXBackend.available(), bool)
    assert isinstance(TransformersBackend.available(), bool)


def test_transformers_load_does_not_blank_cuda_visible_devices():
    """Regression: load() must not hide GPUs via CUDA_VISIBLE_DEVICES=\"\"."""
    import inspect

    src = inspect.getsource(TransformersBackend.load)
    assert 'setdefault("CUDA_VISIBLE_DEVICES"' not in src
    assert "CUDA_VISIBLE_DEVICES\", \"\"" not in src
    assert "CUDA_VISIBLE_DEVICES', ''" not in src


def test_transformers_info_keeps_device_when_name_lookup_fails(monkeypatch):
    """detail is informational; get_device_name failure must not set device=error."""
    import sys
    import types

    be = TransformersBackend(model_id="test/model", device="cuda")
    monkeypatch.setattr(TransformersBackend, "available", staticmethod(lambda: True))
    monkeypatch.setattr(be, "_resolve_device", lambda: "cuda")

    def boom(_i: int) -> str:
        raise RuntimeError("name unavailable")

    fake_cuda = types.SimpleNamespace(
        is_available=lambda: True,
        device_count=lambda: 1,
        get_device_name=boom,
    )
    if "torch" in sys.modules:
        monkeypatch.setattr(sys.modules["torch"], "cuda", fake_cuda)
    else:
        fake_torch = types.ModuleType("torch")
        fake_torch.cuda = fake_cuda  # type: ignore[attr-defined]
        monkeypatch.setitem(sys.modules, "torch", fake_torch)

    info = be.info()
    assert info.device == "cuda"
    assert info.detail == ""
    assert info.available is True


def test_transformers_tensor_device_includes_mps():
    assert TransformersBackend._tensor_device("mps") == "mps"
    assert TransformersBackend._tensor_device("cuda") == "cuda"
    assert TransformersBackend._tensor_device("cpu") == "cpu"
    assert TransformersBackend._tensor_device("unknown") == "cpu"


def test_native_jit_disabled_before_torch_import():
    """torch reads TORCH_DISABLE_NATIVE_JIT at import; the package must set it first."""
    import os
    import subprocess
    import sys

    env = {k: v for k, v in os.environ.items() if k != "TORCH_DISABLE_NATIVE_JIT"}
    code = (
        "import os, sys; import pre_peer_checker; "
        "print('torch' in sys.modules, os.environ.get('TORCH_DISABLE_NATIVE_JIT'))"
    )
    out = subprocess.run(
        [sys.executable, "-c", code], env=env, capture_output=True, text=True, check=True
    ).stdout.split()
    assert out == ["False", "1"]
