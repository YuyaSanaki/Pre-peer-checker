"""Standalone JPEG/PNG figures: panel-letter OCR + letter-anchor regions."""

from __future__ import annotations

from pathlib import Path

from PIL import Image

from pre_peer_checker.parsers.figure_panel_labels import (
    collect_panel_labels_by_figure,
    collect_panel_labels_by_figure_detailed,
    is_publication_figure_raster,
    looks_like_figure_filename,
    panel_labels_from_raster_image_detailed,
)
from pre_peer_checker.parsers.raster_figure_panel_ocr import RasterPanelOcrResult
from pre_peer_checker.report.figure_previews import build_figure_previews, build_raster_figure_preview


def _blank_png(path: Path, size: tuple[int, int] = (80, 60)) -> Path:
    Image.new("RGB", size, "white").save(path)
    return path


def _fake_analysis(_path) -> RasterPanelOcrResult:
    return RasterPanelOcrResult(
        letters={"A", "B"},
        engines=["florence"],
        crops=[
            {"panel": "A", "box": [0.0, 0.0, 40.0, 60.0]},
            {"panel": "B", "box": [40.0, 0.0, 80.0, 60.0]},
        ],
        width=80,
        height=60,
    )


def test_figure_filename_gate():
    assert looks_like_figure_filename(Path("Fig1.png"))
    assert looks_like_figure_filename(Path("FigureS2.jpg"))
    assert is_publication_figure_raster(Path("Fig1.png"))
    assert not is_publication_figure_raster(Path("gel.png"))
    assert not is_publication_figure_raster(Path("Fig1.pdf"))
    assert not looks_like_figure_filename(Path("RplotFig1.png"))


def test_collect_skips_non_figure_png(tmp_path: Path, monkeypatch):
    gel = _blank_png(tmp_path / "gel.png")
    monkeypatch.setattr(
        "pre_peer_checker.parsers.raster_figure_panel_ocr.raster_panel_analysis_from_image",
        _fake_analysis,
    )
    monkeypatch.setattr(
        "pre_peer_checker.parsers.raster_figure_panel_ocr.raster_panel_ocr_enabled",
        lambda: True,
    )
    assert collect_panel_labels_by_figure([gel], raster_fallback=True) == {}


def test_collect_panel_labels_from_png(tmp_path: Path, monkeypatch):
    fig = _blank_png(tmp_path / "Fig1.png")
    monkeypatch.setattr(
        "pre_peer_checker.parsers.raster_figure_panel_ocr.raster_panel_analysis_from_image",
        _fake_analysis,
    )
    by_fig, meta = collect_panel_labels_by_figure_detailed([fig], raster_fallback=True)
    assert by_fig.get("1") == ["A", "B"]
    assert meta["1"].source == "raster_ocr"
    assert meta["1"].needs_review is True
    assert {r["panel"] for r in meta["1"].regions} == {"A", "B"}


def test_photo_boxes_split_edge_to_edge_photos():
    from panel_fixtures import photo

    from pre_peer_checker.parsers.raster_figure_panel_ocr import _photo_boxes

    canvas = Image.new("RGB", (200, 100), "white")
    canvas.paste(photo(71, size=(90, 100)), (0, 0))
    canvas.paste(photo(72, size=(90, 100)), (110, 0))
    boxes = sorted(_photo_boxes(canvas))
    assert len(boxes) == 2
    assert boxes[0][0] == 0 and boxes[0][2] <= 100
    assert boxes[1][0] >= 100 and boxes[1][2] == 200


def test_raster_regions_fit_photos_and_keep_grid_for_graph_panels(tmp_path: Path, monkeypatch):
    """Letters inside or just outside a photo get its box; a photo-less panel keeps its crop."""
    fig = _blank_png(tmp_path / "Fig4.png", (600, 200))

    def analysis(_path) -> RasterPanelOcrResult:
        return RasterPanelOcrResult(
            letters={"A", "B", "C"},
            engines=["florence"],
            crops=[
                {"panel": "A", "box": [0.0, 0.0, 200.0, 200.0]},
                {"panel": "B", "box": [200.0, 0.0, 400.0, 200.0]},
                {"panel": "C", "box": [400.0, 0.0, 600.0, 200.0]},
            ],
            width=600,
            height=200,
            label_boxes={
                "A": [2.0, 5.0, 14.0, 22.0],
                "B": [215.0, 35.0, 228.0, 52.0],
                "C": [405.0, 5.0, 417.0, 22.0],
            },
            photo_boxes=[[20.0, 30.0, 190.0, 190.0], [210.0, 30.0, 380.0, 190.0]],
        )

    monkeypatch.setattr(
        "pre_peer_checker.parsers.raster_figure_panel_ocr.raster_panel_analysis_from_image",
        analysis,
    )
    _labels, meta = panel_labels_from_raster_image_detailed(fig, raster_fallback=True)
    boxes = {r["panel"]: (r["x0"], r["y0"], r["x1"], r["y1"]) for r in meta.regions}
    assert boxes == {
        "A": (2.0, 5.0, 190.0, 190.0),
        "B": (210.0, 30.0, 380.0, 190.0),
        "C": (400.0, 0.0, 600.0, 200.0),
    }


def test_raster_ocr_respects_disabled_flag(tmp_path: Path, monkeypatch):
    fig = _blank_png(tmp_path / "Fig2.png")

    def boom(_path):
        raise AssertionError("OCR must not run when raster_fallback is False")

    monkeypatch.setattr(
        "pre_peer_checker.parsers.raster_figure_panel_ocr.raster_panel_analysis_from_image",
        boom,
    )
    labels, meta = panel_labels_from_raster_image_detailed(fig, raster_fallback=False)
    assert labels == []
    assert meta.source == "none"


def test_raster_preview_from_filename(tmp_path: Path):
    fig = _blank_png(tmp_path / "Fig3.png", (120, 80))
    regions = [
        {
            "panel": "A",
            "source": str(fig),
            "pct": {"left_pct": 0.0, "top_pct": 0.0, "width_pct": 50.0, "height_pct": 100.0},
        }
    ]
    prev = build_raster_figure_preview(fig, regions=regions)
    assert prev is not None
    assert prev["figure_id"] == "Figure 3"
    assert prev["n_panels"] == 1
    assert prev["image_data_uri"].startswith("data:image/jpeg;base64,")
    previews = build_figure_previews([fig], regions=regions)
    assert previews and previews[0]["figure_id"] == "Figure 3"


def test_run_verification_reads_png_panel_labels(tmp_path: Path, monkeypatch):
    from docx import Document

    from pre_peer_checker.pipeline.orchestrator import run_verification

    ms = tmp_path / "case"
    ms.mkdir()
    _blank_png(ms / "Fig1.png")
    doc = Document()
    doc.add_paragraph("Results")
    doc.add_paragraph("Quantification is shown in Fig. 1A and Fig. 1B.")
    doc.add_paragraph("Figure 1. (A) Control. (B) Mutant. n=3.")
    doc.save(ms / "main.docx")

    monkeypatch.setattr(
        "pre_peer_checker.parsers.raster_figure_panel_ocr.raster_panel_analysis_from_image",
        _fake_analysis,
    )
    monkeypatch.setattr(
        "pre_peer_checker.parsers.raster_figure_panel_ocr.raster_panel_ocr_enabled",
        lambda: True,
    )
    result = run_verification([ms], legend_llm=False)
    assert result.artifacts["figure_panel_labels"].get("1") == ["A", "B"]
    assert result.artifacts["figure_panel_labels_meta"]["1"]["source"] == "raster_ocr"
    regions = result.artifacts.get("figure_panel_regions") or []
    assert {r["panel"] for r in regions} >= {"A", "B"}
    chunks = result.artifacts.get("figure_chunks") or []
    assert any(c.get("panel_labels_from_figure") == ["A", "B"] for c in chunks)
    sources = result.artifacts.get("figure_preview_sources") or []
    assert any(str(s).endswith("Fig1.png") for s in sources)


def test_prepare_panel_sources_uses_letter_boxes_when_xy_cut_cannot(tmp_path: Path):
    from panel_fixtures import photo
    from PIL import Image as PILImage

    from pre_peer_checker.imaging.panel_split import PanelBox, split_panels
    from pre_peer_checker.imaging.panel_units import prepare_panel_sources

    a = photo(61, size=(80, 80))
    b = photo(62, size=(80, 80))
    canvas = PILImage.new("RGB", (160, 80))
    canvas.paste(a, (0, 0))
    canvas.paste(b, (80, 0))
    path = tmp_path / "Fig9.png"
    canvas.save(path)
    assert split_panels(PILImage.open(path)) == []
    boxes = [PanelBox(0, 0, 80, 80), PanelBox(80, 0, 160, 80)]
    _prev, units = prepare_panel_sources(
        [path],
        tmp_path / "u",
        n_preview=0,
        panels=True,
        panel_boxes={str(path.resolve()): boxes},
    )
    cropped = [u for u in units if u.box is not None]
    assert len(cropped) == 2
