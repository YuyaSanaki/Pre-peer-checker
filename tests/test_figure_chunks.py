"""Figure-chunked manuscript splitting and check-item extraction."""

from __future__ import annotations

import json
from pathlib import Path

from pre_peer_checker.llm.legend_extract import (
    extract_check_items_from_chunk,
    extract_legends_json_from_docx,
    extract_legends_with_backend,
    legend_json_to_panel_ns,
    merge_panel_ns,
)
from pre_peer_checker.llm.legend_schema import build_figure_chunk_llm_prompt
from pre_peer_checker.parsers.figure_chunks import (
    FigureChunk,
    build_figure_chunks_from_docx,
    build_figure_chunks_from_paragraphs,
    figure_nums_in_text,
)
from pre_peer_checker.parsers.legend_struct import PanelN
from pre_peer_checker.pipeline.orchestrator import run_verification


def test_figure_nums_in_text():
    assert "1" in figure_nums_in_text("as shown in Fig. 1A")
    assert "2" in figure_nums_in_text("see Figure 2")
    nums = figure_nums_in_text("Figs. 1–3 show quantification")
    assert nums >= {"1", "2", "3"}


def test_build_chunks_splits_legend_results_methods():
    paragraphs = [
        "Results",
        "We first examined clone size (Fig. 1A).",
        "producer neurons were also quantified (Figure 1Q).",
        "Methods",
        "Statistical analysis",
        "Welch's t-test was used. Sample sizes are reported in legends.",
        "Imaging for Figure 1 used confocal microscopy.",
        "Figure Legends",
        "Figure 1. Clone size. n=11 (F) and n=17 (N). Welch's t-test.",
        "Figure 2. Control. n=5 (A).",
    ]
    chunks = build_figure_chunks_from_paragraphs(paragraphs)
    by_num = {c.figure_num: c for c in chunks}
    assert "1" in by_num
    assert "2" in by_num
    c1 = by_num["1"]
    assert "n=11 (F)" in c1.legend
    assert any("Fig. 1A" in r or "Figure 1Q" in r for r in c1.results)
    assert c1.methods_includes_shared or any("Welch" in m for m in c1.methods)
    assert "n=5 (A)" in by_num["2"].legend
    # Fig 2 should not pull Fig 1 results
    assert not any("Fig. 1A" in r for r in by_num["2"].results)


def test_chunk_prompt_contains_sections():
    chunk = FigureChunk(
        figure_id="Figure 1",
        figure_num="1",
        legend="Figure 1. n=12 (F).",
        results=["Results mention Fig. 1."],
        methods=["Methods shared stats."],
        methods_includes_shared=True,
    )
    body = chunk.prompt_body()
    assert "[legend]" in body
    assert "[results_mentions]" in body
    assert "[methods_related" in body
    prompt = build_figure_chunk_llm_prompt(body, figure_hint="Figure 1")
    assert "EXTRACTION ONLY" in prompt or "check-item" in prompt.lower()
    assert "n=12 (F)" in prompt


def test_extract_check_items_chunk_with_stub_llm():
    def fake(prompt: str) -> str:
        assert "[legend]" in prompt or "Legend text" in prompt
        return json.dumps(
            {
                "figure": "Figure 1",
                "panels": [
                    {
                        "panel": "F",
                        "n": 12,
                        "groups": ["WT"],
                        "notes": "",
                        "n_scope": "per_group",
                        "evidence_span": "12 biologically independent samples",
                        "confidence": 0.9,
                    }
                ],
                "tests": ["Welch's t-test"],
                "p_values": [],
                "citation": {"mentioned": False, "reproduced_from": None, "spans": []},
                "raw_excerpt": "x",
                "extractor": "llm",
            }
        )

    # Legend uses phrasing rules miss; LLM supplies n
    chunk = FigureChunk(
        figure_id="Figure 1",
        figure_num="1",
        legend="Figure 1. Panel F used 12 biologically independent samples (F).",
        results=["Fig. 1 shows clone ratios."],
        methods=["Statistics used Welch's t-test."],
    )
    rules = extract_check_items_from_chunk(chunk, prefer_llm=False)
    hybrid = extract_check_items_from_chunk(
        chunk, prefer_llm=True, llm_generate=fake
    )
    assert hybrid.extractor == "llm"
    assert any(p.n == 12 and p.panel == "F" for p in hybrid.panels)
    assert any(p.n_scope == "per_group" for p in hybrid.panels)
    # rules may or may not catch this phrasing
    _ = rules


def test_merge_panel_ns_prefers_llm():
    rules = [PanelN(panel="F", n=11, figure="Figure 1", context="n=11")]
    llm = [PanelN(panel="F", n=12, figure="Figure 1", context="n=12")]
    merged = merge_panel_ns(rules, llm)
    assert len(merged) == 1
    assert merged[0].n == 12


def test_docx_chunks_and_orchestrator_artifact(tmp_path: Path):
    from docx import Document

    ms = tmp_path / "ms"
    ms.mkdir()
    doc = Document()
    doc.add_paragraph("Results")
    doc.add_paragraph("Quantification is shown in Fig. 1F.")
    doc.add_paragraph("Methods")
    doc.add_paragraph("Statistical analysis used Welch's t-test.")
    doc.add_paragraph("Figure 1. Clone size. n=5 (A1). Welch's t-test.")
    doc.save(ms / "main.docx")

    chunks = build_figure_chunks_from_docx(ms / "main.docx")
    assert chunks
    assert chunks[0].legend
    assert chunks[0].results

    items, ch2 = extract_legends_json_from_docx(ms / "main.docx")
    assert items and items[0].extractor == "rules"
    assert ch2

    result = run_verification([ms], legend_llm=True, legend_llm_prefer="none")
    assert result.artifacts.get("figure_chunks")
    assert result.artifacts.get("legend_json")
    meta0 = result.artifacts["legend_llm"][0]
    assert meta0.get("figure_chunk_mode") is True
    pns = legend_json_to_panel_ns(items)
    assert any(p.panel == "A1" and p.n == 5 for p in pns)


def test_extract_legends_with_backend_chunk_mode(tmp_path: Path):
    from docx import Document

    docx = tmp_path / "m.docx"
    doc = Document()
    doc.add_paragraph("Results")
    doc.add_paragraph("See Fig. 1.")
    doc.add_paragraph("Figure 1. n=6 (A).")
    doc.save(docx)
    legs, meta = extract_legends_with_backend(docx, enabled=False)
    assert meta["figure_chunk_mode"] is True
    assert meta["n_figure_chunks"] >= 1
    assert legs
