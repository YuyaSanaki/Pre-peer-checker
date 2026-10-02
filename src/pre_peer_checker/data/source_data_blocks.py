"""Journal "Source Data" workbooks: one table block per figure panel.

Layouts seen in Nature / Springer ESM, Cell, eLife and lab exports::

    Fig. 4h                                  Fig. 4l
    scrib vs wild-type (Fig. 4b)             wild-type vs wild-type (Fig. 4i)
    n | Area clone | Area disc | Rel. size   clone no | Perimeter | cCasp3 | ratio
    1 | 47329      | 336005    | 14.08      1        | 395.6     | 1      | 0.0025
    ...
                     Mean      | 11.27
                     SD        | 2.13

    sheet "2b" of Source_Data_Fig2.xlsx      sheet "Fig.1c"
    media      | time | primedScore          Fig.1c
    Fibroblast | D0   | 0.107                Fecundity
    Primed     | D13  | 0.462                Ctrl  | MetR
                                             0.92  | 0.04

Each header row + the rows below it becomes one :class:`SourceDataBlock`. Text
columns in the body (``media``, ``time``) are categories: rows sharing them form
one group (tidy / long format). Without categories and without a sample-index
column, each numeric column is one group (wide format).

The figure label comes from a title cell above the block, else the sheet name
(``Fig.1c``, ``2b``, ``3bf``), else the file name (``Source_Data_ED_Fig3``). The
reference in a title (``Fig. 4b``) names the image panel of the condition.
"""

from __future__ import annotations

import re
from collections import OrderedDict
from dataclasses import dataclass, field, replace
from functools import cached_property, lru_cache
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter

from pre_peer_checker.data.curve_n import alive_count_cohort, km_cohort_size
from pre_peer_checker.data.figure_refs import (
    FigureLabel,
    figure_from_filename,
    looks_like_source_data_file,
    parse_figure_label,
    sheet_figure_label,
    workbook_uses_short_sheet_names,
)
from pre_peer_checker.data.table_grid import cell_kind

_FIG_REF_RE = re.compile(
    r"\b(?:(?P<ext>extended\s+data|supplementary)\s+)?"
    r"fig(?:ure)?s?\.?\s*(?P<num>S?\d+)\s*(?P<panel>[a-z])?\b",
    re.IGNORECASE,
)
_SUMMARY_RE = re.compile(
    r"^\s*(mean|average|ave\.?|avg\.?|sd|s\.d\.?|stdev|std\.?\s*dev\.?|sem|s\.e\.m\.?|"
    r"se|s\.e\.?|median|n|count|sum|total|max|min|variance|var)\s*:?\s*$",
    re.IGNORECASE,
)
_INDEX_HEADER_RE = re.compile(
    r"^\s*(n|no\.?|#|id|sample(\s*(no\.?|id))?|clone\s*(no\.?|#|id)|disc\s*(no\.?|id)|"
    r"dic\s*id|cell\s*(no\.?|#|id)|animal(\s*id)?|mouse(\s*id)?|replicate|rep\.?|"
    r"embryo(\s*(no\.?|id))?|larva(\s*(no\.?|id))?)\s*$",
    re.IGNORECASE,
)
# x-axis of a curve (survival, time course, dose response): rows are not samples
_XAXIS_HEADER_RE = re.compile(
    r"^\s*(days?|d|hours?|hrs?|h|min(ute)?s?|sec(ond)?s?|s|time(\s*\(.*\))?|t|age|weeks?|"
    r"wk|months?|x|dose|conc\.?|concentration|distance|position|frame|cycle|"
    r"wavelength|radius)\s*(\(.*\))?\s*$",
    re.IGNORECASE,
)

_SUMMARY_CANON = {
    "mean": "mean", "average": "mean", "ave": "mean", "ave.": "mean", "avg": "mean",
    "avg.": "mean", "sd": "sd", "s.d.": "sd", "s.d": "sd", "stdev": "sd",
    "sem": "sem", "s.e.m.": "sem", "s.e.m": "sem", "se": "sem", "s.e.": "sem",
    "s.e": "sem", "median": "median", "n": "n", "count": "n",
}

# Blocks with more rows than this are per-cell / per-feature dumps (single-cell
# coordinates, omics tables), not per-sample values a legend n counts.
LARGE_BLOCK_ROWS = 1000

# "Experiment 2" / "Rep 3" / "Batch 1" rows split one group's column into runs
_REPLICATE_RE = re.compile(
    r"^\s*(?:(?:independent\s+|biological\s+|technical\s+)?(?:experiment|replicate)s?|"
    r"exp\.?|expt\.?|rep\.?|batch|trial|run|round|set|cohort|plate|donor|litter)"
    r"\s*[#no.]*\s*\d{1,3}\s*$",
    re.IGNORECASE,
)

# survival-type curves (cohort recoverable) vs other time courses (no sample n)
_SURVIVAL_RE = re.compile(
    r"surviv|life\s*span|longevity|mortality|death|dead|alive|kaplan|resistan",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class FigureRef:
    number: str  # "4", "S2"
    panel: str  # "h" or ""
    extended: bool = False

    def label(self) -> str:
        prefix = "Extended Data Fig. " if self.extended else "Fig. "
        return f"{prefix}{self.number}{self.panel}"


@dataclass
class SourceColumn:
    header: str
    values: list[float | None]
    summary: dict[str, float] = field(default_factory=dict)

    def numeric(self) -> list[float]:
        return [v for v in self.values if v is not None]


@dataclass
class CategoryColumn:
    header: str
    values: list[str | None]

    @property
    def id_like(self) -> bool:
        """Every row has its own label (sample / gene names), so it does not group rows."""
        vals = [v for v in self.values if v is not None]
        return len(vals) >= 3 and len(set(vals)) == len(vals)


@dataclass
class SourceDataBlock:
    path: Path
    sheet: str
    header_cell: str
    figure: FigureRef | None
    title: str
    panel_refs: tuple[FigureRef, ...]
    columns: list[SourceColumn]
    has_index: bool
    panels: tuple[str, ...] = ()
    categories: list[CategoryColumn] = field(default_factory=list)
    truncated: bool = False
    # y columns of a curve whose x column sits in a separate block to the left
    curve: bool = False
    # figure taken from a bare sheet name like "3bf" (main vs Extended Data unknown)
    figure_implicit: bool = False

    @property
    def primary(self) -> SourceColumn:
        """Plotted quantity: the column carrying Mean/SD rows, else the last one."""
        with_summary = [c for c in self.columns if c.summary.get("mean") is not None]
        return (with_summary or self.columns)[-1]

    @property
    def n(self) -> int:
        return len(self.primary.numeric())

    @property
    def label(self) -> str:
        parts = []
        if self.figure is not None:
            fig = self.figure.label()
            if len(self.panels) > 1:
                fig = fig[: len(fig) - len(self.figure.panel)] + "".join(self.panels)
            parts.append(fig)
        if self.title:
            parts.append(self.title)
        return " | ".join(parts) or f"{self.sheet}!{self.header_cell}"

    @property
    def location(self) -> str:
        return f"{self.path.name} [{self.sheet}!{self.header_cell}]"

    @property
    def all_panels(self) -> tuple[str, ...]:
        if self.panels:
            return self.panels
        if self.figure is not None and self.figure.panel:
            return (self.figure.panel,)
        return ()

    @property
    def grouping_categories(self) -> list[CategoryColumn]:
        return [c for c in self.categories if not c.id_like]

    @property
    def layout(self) -> str:
        """``long`` | ``wide`` | ``single`` (rows are samples) or ``xy`` | ``matrix`` | ``large``."""
        rows = max((len(c.values) for c in self.columns), default=0)
        if self.truncated or rows > LARGE_BLOCK_ROWS:
            return "large"
        if self.categories:
            if self.grouping_categories:
                return "long"
            return "matrix" if len(self.columns) >= 3 else "single"
        if self.curve or (len(self.columns) == 1 and _is_x_axis(self.columns[0])):
            return "xy"
        if len(self.columns) >= 2 and _is_x_axis(self.columns[0]):
            return "xy"
        if (
            len(self.columns) >= 2
            and not self.has_index
            and not any(c.summary for c in self.columns)
        ):
            return "wide"
        return "single"

    @property
    def n_comparable(self) -> bool:
        """Whether the block yields a sample n a legend n can be checked against."""
        return self.layout in {"long", "wide", "single"} or self.curve_ns is not None

    @cached_property
    def curve_ns(self) -> list[tuple[str, int, bool]] | None:
        """Survival curves: (group, cohort n, exact) when every curve reconstructs.

        Kaplan–Meier percentages give the number at risk at the first death, a
        lower bound (animals censored earlier leave the curve unchanged); columns
        of animals alive per time point give n exactly.
        """
        if self.layout != "xy":
            return None
        context = " ".join([self.title, self.sheet, *(c.header or "" for c in self.columns)])
        if not _SURVIVAL_RE.search(context):
            return None
        ys = self.columns if self.curve else self.columns[1:]
        out: list[tuple[str, int, bool]] = []
        for c in ys:
            vals = c.numeric()
            if not vals:
                continue
            km = km_cohort_size(vals) if max(vals) <= 100.0 + 1e-9 else None
            if km is not None:
                out.append((c.header, km, False))
                continue
            alive = alive_count_cohort(vals)
            if alive is None:
                return None
            out.append((c.header, alive, True))
        return out or None

    @property
    def group_ns(self) -> list[tuple[str, int, bool]]:
        """(group, n, exact) — row counts for sample tables, cohorts for survival curves."""
        curves = self.curve_ns
        if curves is not None:
            return curves
        return [(g, len(v), True) for g, v in self.groups]

    def group_n(self, group: str) -> tuple[int, bool] | None:
        hit = _match_group([(g, []) for g, _, _ in self.group_ns], group)
        if hit is None:
            return None
        return next((n, ex) for g, n, ex in self.group_ns if g == hit[0])

    @property
    def groups(self) -> list[tuple[str, list[float]]]:
        """(group name, values) — categories (long), columns (wide / xy), else primary."""
        layout = self.layout
        if layout == "long":
            cats = self.grouping_categories
            col = self.primary
            out: OrderedDict[str, list[float]] = OrderedDict()
            for i, v in enumerate(col.values):
                key = " / ".join(c.values[i] or "" for c in cats).strip(" /")
                out.setdefault(key or "all", [])
                if v is not None:
                    out[key or "all"].append(v)
            return [(k, vals) for k, vals in out.items() if vals]
        if layout == "wide":
            return [(c.header, c.numeric()) for c in self.columns if c.numeric()]
        if layout == "xy":
            ys = self.columns if self.curve else self.columns[1:]
            return [(c.header, c.numeric()) for c in ys if c.numeric()]
        return [(self.primary.header or "all", self.primary.numeric())]

    def group_key(self, group: str) -> str:
        """GroupVector key: the block label, plus the group when there are several."""
        return self.label if len(self.groups) <= 1 else f"{self.label} | {group}"

    def legend_n_conflict(self, legend_n: int, group: str = "") -> int | None:
        """Row count that disagrees with ``legend_n`` (None when consistent / not comparable)."""
        if not self.n_comparable:
            return None
        if group:
            hit = self.group_n(group)
            if hit is not None:
                return None if _n_agrees(legend_n, *hit) else hit[0]
        ns = [(n, ex) for _, n, ex in self.group_ns] or [(self.n, True)]
        bad = [n for n, ex in ns if not _n_agrees(legend_n, n, ex)]
        return bad[0] if bad else None


def _n_agrees(legend_n: int, data_n: int, exact: bool) -> bool:
    return legend_n == data_n if exact else legend_n >= data_n


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9+]", "", s.lower())


def _match_group(
    groups: list[tuple[str, list[float]]], name: str
) -> tuple[str, list[float]] | None:
    key = _norm(name)
    if not key:
        return None
    exact = [g for g in groups if _norm(g[0]) == key]
    if len(exact) == 1:
        return exact[0]
    if len(key) < 3:
        return None
    part = [g for g in groups if key in _norm(g[0]) or (len(_norm(g[0])) >= 3 and _norm(g[0]) in key)]
    return part[0] if len(part) == 1 else None


def match_group(block: SourceDataBlock, name: str) -> tuple[str, list[float]] | None:
    """Group of ``block`` named like a legend group (``Ctrl``, ``MetR``)."""
    return _match_group(block.groups, name)


def content_signature(block: SourceDataBlock) -> tuple:
    """Equal for copies of the same table (duplicated files / sheets)."""
    return (
        block.figure,
        block.all_panels,
        tuple((g, tuple(round(v, 9) for v in vals)) for g, vals in block.groups),
    )


LAYOUT_JA = {
    "xy": "時系列・曲線データ（行は時点）",
    "matrix": "行列データ（行は遺伝子・項目）",
    "large": "大規模表（行は細胞・特徴量）",
}


def _is_x_axis(col: SourceColumn) -> bool:
    """Monotonic first column named like time / dose / distance."""
    vals = col.numeric()
    if len(vals) < 3 or len(vals) != len(col.values):
        return False
    increasing = all(b > a for a, b in zip(vals, vals[1:]))
    return increasing and bool(_XAXIS_HEADER_RE.match(col.header or ""))


def _cell_kind(v: Any) -> tuple[str, Any]:
    return cell_kind(v)


def _read_grids(path: Path) -> list[tuple[str, list[list[tuple[str, Any]]], bool]]:
    from pre_peer_checker.data.table_grid import read_grids

    return [(g.sheet, g.rows, g.truncated) for g in read_grids(path)]


def _cell(grid: list[list[tuple[str, Any]]], r: int, c: int) -> tuple[str, Any]:
    if r < 0 or r >= len(grid):
        return "empty", None
    row = grid[r]
    if c < 0 or c >= len(row):
        return "empty", None
    return row[c]


def _is_summary(kind: str, v: Any) -> bool:
    return kind == "str" and bool(_SUMMARY_RE.match(v))


def _summary_key(v: str) -> str | None:
    return _SUMMARY_CANON.get(v.strip().rstrip(":").strip().lower())


def _as_ref(lbl: FigureLabel) -> FigureRef:
    return FigureRef(number=lbl.number, panel=lbl.panel, extended=lbl.extended)


def _parse_figure_label(s: str) -> FigureRef | None:
    lbl = parse_figure_label(s)
    return _as_ref(lbl) if lbl is not None else None


def _figure_refs(s: str) -> tuple[FigureRef, ...]:
    return tuple(
        FigureRef(
            number=m.group("num").upper(),
            panel=(m.group("panel") or "").lower(),
            extended=bool(m.group("ext")),
        )
        for m in _FIG_REF_RE.finditer(s)
    )


def _fmt_category(kind: str, v: Any) -> str | None:
    if kind == "empty":
        return None
    if kind == "num":
        return str(int(v)) if float(v).is_integer() else str(v)
    return str(v)


def _numeric_header_labels(grid: list[list[tuple[str, Any]]]) -> list[list[tuple[str, Any]]]:
    """Header rows where a label such as ``100%`` was stored as the number 100.

    A number flanked by string headers that both head numeric columns, with a
    non-numeric cell above it, is a column label rather than a data value.
    """
    out = grid
    for r, row in enumerate(grid):
        if sum(1 for k, _ in row if k == "str") < 2:
            continue
        for c in range(1, len(row) - 1):
            if row[c][0] != "num" or row[c - 1][0] != "str" or row[c + 1][0] != "str":
                continue
            if _cell(grid, r - 1, c)[0] == "num":
                continue
            if not all(_cell(grid, r + 1, j)[0] == "num" for j in (c - 1, c, c + 1)):
                continue
            if out is grid:
                out = list(grid)
            if out[r] is row:
                out[r] = list(row)
            v = row[c][1]
            out[r][c] = ("str", f"{v:g}" if isinstance(v, float) else str(v))
    return out


def _header_runs(grid: list[list[tuple[str, Any]]], r: int) -> list[tuple[int, int]]:
    """Contiguous string runs in row ``r`` that sit on top of data rows."""
    row = grid[r]
    runs: list[tuple[int, int]] = []
    c = 0
    while c < len(row):
        if row[c][0] != "str":
            c += 1
            continue
        start = c
        while c < len(row) and row[c][0] == "str":
            c += 1
        end = c - 1
        cells = [row[j][1] for j in range(start, end + 1)]
        if all(_SUMMARY_RE.match(x) and not _INDEX_HEADER_RE.match(x) for x in cells):
            continue
        if end - start + 1 == 1 and _parse_figure_label(cells[0]):
            continue
        if all(_REPLICATE_RE.match(x) for x in cells) and any(
            _cell(grid, r - 1, j)[0] == "num" for j in range(start, end + 1)
        ):
            continue
        below = [_cell(grid, r + 1, j)[0] for j in range(start, end + 1)]
        n_num = below.count("num")
        if n_num >= 1 and (n_num * 2 >= len(below) or n_num + below.count("str") == len(below)):
            runs.append((start, end))
    return runs


def _nearest_figure_label(
    grid: list[list[tuple[str, Any]]], r: int, c0: int, c1: int
) -> FigureLabel | None:
    for rr in range(r - 1, -1, -1):
        for cc in range(c0, c1 + 1):
            k, v = _cell(grid, rr, cc)
            if k == "str":
                ref = parse_figure_label(v)
                if ref is not None:
                    return ref
    best: tuple[int, int, FigureLabel] | None = None
    for rr in range(r - 1, -1, -1):
        for cc in range(c0):
            k, v = _cell(grid, rr, cc)
            if k != "str":
                continue
            ref = parse_figure_label(v)
            if ref is None:
                continue
            score = (r - rr, c0 - cc)
            if best is None or score < best[:2]:
                best = (*score, ref)
    return best[2] if best else None


def _is_replicate_row(cells: list[tuple[str, Any]]) -> bool:
    filled = [(k, v) for k, v in cells if k != "empty"]
    return bool(filled) and all(k == "str" and _REPLICATE_RE.match(v) for k, v in filled)


def _stacked_header(grid: list[list[tuple[str, Any]]], r: int, c: int, c0: int) -> list[str]:
    """Header cells above (r, c); a merged cell left of c within the run counts too."""
    parts: list[str] = []
    for up in range(r - 1, max(-1, r - 4), -1):
        k, v = _cell(grid, up, c)
        if k == "empty":
            left = next(
                (_cell(grid, up, cc) for cc in range(c - 1, c0 - 1, -1) if _cell(grid, up, cc)[0] != "empty"),
                ("empty", None),
            )
            k, v = left
        if k != "str" or parse_figure_label(v) is not None or _is_summary(k, v):
            break
        parts.insert(0, str(v).strip())
    return parts


def _resolve_headers(grid: list[list[tuple[str, Any]]], r: int, c0: int, c1: int) -> list[str]:
    """Column names; stacked/merged header rows name groups whose own cell is a replicate
    label ("Experiment 1") or repeats a neighbour ("shX", "shX")."""
    raw = [str(_cell(grid, r, c)[1]) for c in range(c0, c1 + 1)]
    weak = [bool(_REPLICATE_RE.match(h)) for h in raw]
    dup = len(set(raw)) < len(raw)
    if not any(weak) and not dup:
        return raw
    out: list[str] = []
    for i, c in enumerate(range(c0, c1 + 1)):
        stack = _stacked_header(grid, r, c, c0)
        own = [] if weak[i] else [raw[i]]
        out.append(" ".join(stack + own) or raw[i])
    counts = {h: out.count(h) for h in out}
    return [
        f"{h} [{get_column_letter(c0 + i + 1)}]" if counts[h] > 1 else h
        for i, h in enumerate(out)
    ]


def _parse_block(
    path: Path,
    sheet: str,
    grid: list[list[tuple[str, Any]]],
    r: int,
    c0: int,
    c1: int,
    *,
    sheet_label: FigureLabel | None = None,
    sheet_label_implicit: bool = False,
    truncated: bool = False,
) -> SourceDataBlock | None:
    width = c1 - c0 + 1
    headers = _resolve_headers(grid, r, c0, c1)
    first = [_cell(grid, r + 1, c) for c in range(c0, c1 + 1)]
    cat_idx = {i for i, (k, v) in enumerate(first) if k == "str" and not _is_summary(k, v)}
    if len(cat_idx) == width:
        return None
    nums: list[SourceColumn] = [SourceColumn(header=h, values=[]) for h in headers]
    cats: dict[int, list[str | None]] = {i: [] for i in cat_idx}

    rr = r + 1
    n_data = 0
    while rr < len(grid):
        cells = [_cell(grid, rr, c) for c in range(c0, c1 + 1)]
        if all(k == "empty" for k, _ in cells):
            break
        summary_at = next(
            (i for i, (k, v) in enumerate(cells) if _is_summary(k, v)), None
        )
        if summary_at is not None and not cat_idx:
            key = _summary_key(cells[summary_at][1])
            if key is not None:
                for i in range(summary_at + 1, len(cells)):
                    k, v = cells[i]
                    if k == "num":
                        nums[i].summary.setdefault(key, v)
            rr += 1
            continue
        if not cat_idx and _is_replicate_row(cells):
            rr += 1
            continue
        if any(k == "str" and i not in cat_idx for i, (k, _) in enumerate(cells)):
            break
        if any(c.summary for c in nums):
            break
        if cat_idx and all(cells[i][0] == "empty" for i in range(width) if i not in cat_idx):
            break
        for i, (k, v) in enumerate(cells):
            if i in cat_idx:
                cats[i].append(_fmt_category(k, v))
            else:
                nums[i].values.append(v if k == "num" else None)
        n_data += 1
        rr += 1
    if n_data < 2:
        return None

    num_cols = [nums[i] for i in range(width) if i not in cat_idx]
    has_index = False
    if len(num_cols) >= 2 and 0 not in cat_idx:
        first_col = nums[0]
        seq = first_col.numeric()
        is_seq = len(seq) >= 2 and seq == [float(i) for i in range(1, len(seq) + 1)]
        if _INDEX_HEADER_RE.match(first_col.header) or (is_seq and not first_col.summary):
            has_index = True
            num_cols = num_cols[1:]
    num_cols = [c for c in num_cols if c.numeric()]
    if not num_cols:
        return None
    categories = [CategoryColumn(header=headers[i], values=cats[i]) for i in sorted(cat_idx)]

    title_parts: list[str] = []
    label: FigureLabel | None = None
    for up in range(r - 1, max(-1, r - 5), -1):
        cells = [_cell(grid, up, c) for c in range(c0, c1 + 1)]
        if any(k == "num" for k, _ in cells) or any(_is_summary(k, v) for k, v in cells):
            break
        for k, v in cells:
            if k != "str":
                continue
            ref = parse_figure_label(v)
            if ref is not None:
                label = label or ref
            else:
                title_parts.insert(0, v)
    if label is None:
        label = _nearest_figure_label(grid, r, c0, c1)
    implicit = False
    if label is None:
        label = sheet_label
        implicit = sheet_label_implicit and label is not None
    title = " / ".join(title_parts)

    return SourceDataBlock(
        path=path,
        sheet=sheet,
        header_cell=f"{get_column_letter(c0 + 1)}{r + 1}",
        figure=_as_ref(label) if label is not None else None,
        title=title,
        panel_refs=_figure_refs(title),
        columns=num_cols,
        has_index=has_index,
        panels=label.panels if label is not None else (),
        categories=categories,
        truncated=truncated,
        figure_implicit=implicit,
    )


def _file_label(path: Path) -> FigureLabel | None:
    """Figure of a whole file — only for Source Data style names, not lab file names."""
    lbl = figure_from_filename(path)
    if lbl is None:
        return None
    if looks_like_source_data_file(path) or parse_figure_label(path.stem) is not None:
        return lbl
    return None


def _parse_uncached(path: Path) -> list[SourceDataBlock]:
    try:
        grids = _read_grids(path)
    except Exception:  # noqa: BLE001
        return []
    file_label = _file_label(path)
    short_ok = file_label is not None or workbook_uses_short_sheet_names(
        [name for name, _, _ in grids], series_guard=not looks_like_source_data_file(path)
    )
    blocks: list[SourceDataBlock] = []
    for sheet, grid, truncated in grids:
        sheet_label = sheet_figure_label(sheet, file_label=file_label, short_names_ok=short_ok)
        if sheet_label is None:
            sheet_label = file_label
        implicit = file_label is None and parse_figure_label(sheet) is None
        grid = _numeric_header_labels(grid)
        for r in range(len(grid)):
            prev: SourceDataBlock | None = None
            for c0, c1 in _header_runs(grid, r):
                block = _parse_block(
                    path,
                    sheet,
                    grid,
                    r,
                    c0,
                    c1,
                    sheet_label=sheet_label,
                    sheet_label_implicit=implicit,
                    truncated=truncated,
                )
                if block is None:
                    continue
                if (
                    prev is not None
                    and len(prev.columns) == 1
                    and not prev.categories
                    and _is_x_axis(prev.columns[0])
                    and not block.categories
                    and max(len(c.values) for c in block.columns) <= len(prev.columns[0].values)
                ):
                    block.curve = True
                blocks.append(block)
                prev = block
    if not any(b.figure is not None for b in blocks):
        return []
    return blocks


@lru_cache(maxsize=256)
def _parse_cached(path_str: str, mtime_ns: int) -> tuple[SourceDataBlock, ...]:
    return tuple(_parse_uncached(Path(path_str)))


def parse_source_data_blocks(path: Path | str) -> list[SourceDataBlock]:
    """Blocks of a Source Data workbook; empty unless a sheet / file names a figure."""
    p = Path(path)
    if p.suffix.lower() not in {".xlsx", ".xlsm", ".xls", ".xlsb", ".ods"}:
        return []
    try:
        st = p.stat()
    except OSError:
        return []
    return list(_parse_cached(str(p.resolve()), st.st_mtime_ns))


def is_source_data_workbook(path: Path | str) -> bool:
    return bool(parse_source_data_blocks(path))


def _values_key(block: SourceDataBlock) -> tuple:
    return tuple((g, tuple(round(v, 9) for v in vals)) for g, vals in block.groups)


def reconcile_copies(blocks: list[SourceDataBlock]) -> list[SourceDataBlock]:
    """A copy labelled only by a bare sheet name takes the figure of an explicit copy.

    Journals often ship the same Source Data twice (e.g. renamed supplementary
    files); the copy whose file / cell says "Extended Data" decides the figure.
    """
    explicit: dict[tuple, FigureRef] = {}
    for b in blocks:
        if b.figure is not None and not b.figure_implicit:
            explicit.setdefault((b.figure.number, b.all_panels, _values_key(b)), b.figure)
    if not explicit:
        return list(blocks)
    out: list[SourceDataBlock] = []
    for b in blocks:
        if b.figure is not None and b.figure_implicit:
            ref = explicit.get((b.figure.number, b.all_panels, _values_key(b)))
            if ref is not None and ref != b.figure:
                b = replace(b, figure=ref, figure_implicit=False)
        out.append(b)
    return out


@lru_cache(maxsize=32)
def _bundle_cached(key: tuple[tuple[str, int], ...]) -> tuple[SourceDataBlock, ...]:
    blocks: list[SourceDataBlock] = []
    for path_str, _ in key:
        blocks.extend(parse_source_data_blocks(path_str))
    return tuple(reconcile_copies(blocks))


def parse_source_data_bundle(paths) -> list[SourceDataBlock]:
    """Blocks of all Source Data workbooks of a submission, duplicates reconciled."""
    key = []
    for p in dict.fromkeys(Path(x).resolve() for x in paths):
        try:
            key.append((str(p), p.stat().st_mtime_ns))
        except OSError:
            continue
    return list(_bundle_cached(tuple(sorted(key))))
