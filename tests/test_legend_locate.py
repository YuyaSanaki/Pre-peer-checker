from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from pre_peer_checker.parsers.figure_chunks import figure_num_key, legend_coverage
from pre_peer_checker.parsers.legend_struct import (
    StructuredLegend,
    extract_structured_legends,
    legend_figure_name,
    match_legend_head,
)
from pre_peer_checker.parsers.manuscript_text import _LEGEND_HEAD_RE, docx_paragraphs
from pre_peer_checker.parsers.references import extract_references_section


def _caption_docx(path: Path) -> Path:
    doc = Document()
    doc.add_paragraph("Results")
    doc.add_paragraph("Figure 2 shows the dose response (Fig. 1).")
    p = doc.add_paragraph("Figure ")
    p._p.append(
        parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr=" SEQ Figure \\* ARABIC ">'
            "<w:r><w:t>1</w:t></w:r></w:fldSimple>"
        )
    )
    p.add_run(". Body weight. n = 8 mice per group.")
    ins = doc.add_paragraph("Data are mean ± SEM ")
    ins._p.append(
        parse_xml(
            f'<w:ins {nsdecls("w")} w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z">'
            "<w:r><w:t>(tracked insertion)</w:t></w:r></w:ins>"
        )
    )
    table = doc.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "[figure image]"
    table.cell(1, 0).text = "Figure 2. Dose response. n = 6 per dose."
    doc.save(path)
    return path


def test_docx_reads_caption_fields_tables_and_insertions(tmp_path: Path) -> None:
    path = _caption_docx(tmp_path / "m.docx")
    paras = docx_paragraphs(path)
    assert "Figure 1. Body weight. n = 8 mice per group." in paras
    assert "Data are mean ± SEM (tracked insertion)" in paras
    assert "Figure 2. Dose response. n = 6 per dose." in paras

    legends = {leg.figure: leg for leg in extract_structured_legends(path)}
    assert set(legends) == {"Figure 1", "Figure 2"}
    assert "tracked insertion" in legends["Figure 1"].text
    assert legends["Figure 2"].text.startswith("Figure 2. Dose response")


def test_pdf_legend_head_styles() -> None:
    for head in (
        "Figure 1. Title",
        "FIGURE 1. Title",
        "Fig. 1 | Title",
        "Figure 1: Title",
        "Figure 1 Loss of function",
        "Supplementary Fig. 2. Title",
        "Extended Data Fig. 3 | Title",
    ):
        assert _LEGEND_HEAD_RE.match(head), head
    for body in ("Figure 1 shows that", "Figure 1—figure supplement 1. Title", "Figure 1"):
        assert not _LEGEND_HEAD_RE.match(body), body


def test_match_legend_head_skips_body_sentences_and_sub_items() -> None:
    assert match_legend_head("Figure 2 shows the dose response.") is None
    assert match_legend_head("Figure 2A shows the dose response.") is None
    assert match_legend_head("Figure 1—figure supplement 1. Extra panels.") is None
    assert match_legend_head("Figure 1 (continued)") is None
    m = match_legend_head("FIGURE 3. Title")
    assert m and legend_figure_name(m.group(1), m.group(2)) == "Figure 3"
    m = match_legend_head("Supplementary Fig. 2. Title")
    assert m and legend_figure_name(m.group(1), m.group(2)) == "Supplementary Figure S2"
    m = match_legend_head("Fig 4 Title")
    assert m and legend_figure_name(m.group(1), m.group(2)) == "Figure 4"


def test_continued_legend_joins_its_figure() -> None:
    from pre_peer_checker.parsers.legend_struct import extract_structured_legends_from_paragraphs

    legends = extract_structured_legends_from_paragraphs(
        [
            "Figure 1. Body weight. (A) Weight curve.",
            "Figure 1 (continued) (B) Food intake. n = 8 mice.",
            "Figure 2. Dose response.",
        ]
    )
    assert [leg.figure for leg in legends] == ["Figure 1", "Figure 2"]
    assert legends[0].text == "Figure 1. Body weight. (A) Weight curve. (B) Food intake. n = 8 mice."


def test_figure_keys_keep_supplementary_and_tables_apart() -> None:
    assert figure_num_key("Supplementary Figure 2") == "S2"
    assert figure_num_key("Supplementary Figure S2") == "S2"
    assert figure_num_key("Table 1") == "T1"
    assert figure_num_key("Figure 1") == "1"
    assert figure_num_key("Extended Data Figure 3") == "ED3"


def test_legend_coverage_reports_missing_figures() -> None:
    paragraphs = [
        "Results",
        "Mice lost weight (Fig. 1a) and responded to dose (Fig. 2b).",
        "Knockout abolished the response (Figure 3c), unlike a prior report (their Fig. 9).",
        "Supplementary Fig. 1 shows controls (Supplementary Figs. 4, 5).",
        "Figure legends",
        "Figure 1. Body weight.",
        "Figure 2. Dose response.",
    ]
    legends = [
        StructuredLegend(figure="Figure 1", text="Figure 1. Body weight."),
        StructuredLegend(figure="Figure 2", text="Figure 2. Dose response."),
    ]
    cov = legend_coverage(paragraphs, legends)
    assert cov["expected"] == ["1", "2", "3"]
    assert cov["missing"] == ["3"]

    cov = legend_coverage(paragraphs, legends, ["1", "2", "3", "4"])
    assert cov["missing"] == ["3", "4"]

    cov = legend_coverage(paragraphs, [])
    assert cov["missing"] == ["1", "2", "3"]


def test_references_stop_at_trailing_captions() -> None:
    paras = [
        "References",
        "1. Smith J. A study. Nature. 2020;1:1-2.",
        "Table 1. Baseline characteristics",
        "Age",
    ]
    assert extract_references_section(paras) == ["1. Smith J. A study. Nature. 2020;1:1-2."]
