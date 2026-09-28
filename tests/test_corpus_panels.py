"""H3 panel-level corpus matching: raster panel split, Word embeds, degraded reuse."""

from __future__ import annotations

import io
from pathlib import Path

import numpy as np
import pytest
from PIL import Image, ImageDraw

from pre_peer_checker.imaging.panel_split import is_photo_like, split_panels
from pre_peer_checker.parsers.docx_images import export_docx_images


def _photo(seed: int, size: tuple[int, int] = (360, 360)) -> Image.Image:
    """1/f colour texture: a stand-in photo whose keypoints survive downscale + JPEG."""
    rng = np.random.default_rng(seed)
    w, h = size
    base = np.zeros((h, w, 3), dtype=np.float32)
    for cells in (3, 6, 12, 24, 48):
        noise = (rng.random((cells, cells, 3)) * 255).astype(np.uint8)
        layer = Image.fromarray(noise).resize((w, h), Image.Resampling.BICUBIC)
        base += np.asarray(layer, dtype=np.float32) / cells**0.5
    base = (base - base.min()) / (base.max() - base.min()) * 255
    return Image.fromarray(base.astype(np.uint8))


def _page(photos: list[Image.Image], *, gap: int = 40, size=(1400, 1000)) -> Image.Image:
    page = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(page)
    draw.text((60, 20), "Figure S1. Composite with plots and text", fill="black")
    x = 60
    for ph in photos:
        page.paste(ph, (x, 120))
        x += ph.size[0] + gap
    # bar-chart-like region: mostly white, must not become a panel
    for i in range(5):
        draw.rectangle((80 + i * 60, 900 - 40 * (i + 1), 110 + i * 60, 900), outline="black")
    return page


def test_split_panels_finds_photos_and_skips_plots():
    photos = [_photo(1), _photo(2), _photo(3)]
    boxes = split_panels(_page(photos))
    assert len(boxes) == 3
    for b in boxes:
        assert abs(b.width - 360) <= 4 and abs(b.height - 360) <= 4


def test_split_panels_single_photo_is_not_split():
    assert split_panels(_photo(4, (600, 450))) == []


def test_split_panels_dense_fallback_when_connector_blocks_cut():
    page = _page([_photo(5), _photo(6)])
    draw = ImageDraw.Draw(page)
    # a connector line running under both photos and up the side blocks every XY cut
    draw.line((20, 100, 20, 980), fill="red", width=6)
    draw.line((20, 980, 1380, 980), fill="red", width=6)
    draw.line((20, 100, 1380, 100), fill="red", width=6)
    boxes = split_panels(page)
    assert len(boxes) == 2


def test_is_photo_like():
    assert is_photo_like(_photo(7))
    assert not is_photo_like(_page([]))


def test_export_docx_images_dedupes_across_versions(tmp_path: Path):
    from docx import Document

    buf = io.BytesIO()
    _photo(8, (300, 200)).save(buf, format="PNG")
    tiny = io.BytesIO()
    Image.new("RGB", (40, 40), "red").save(tiny, format="PNG")
    for name in ("supp_v1.docx", "supp_v2.docx"):
        doc = Document()
        doc.add_picture(io.BytesIO(buf.getvalue()))
        doc.add_picture(io.BytesIO(tiny.getvalue()))
        doc.save(tmp_path / name)
    seen: set[str] = set()
    first = export_docx_images(tmp_path / "supp_v1.docx", tmp_path / "o1", seen_hashes=seen)
    second = export_docx_images(tmp_path / "supp_v2.docx", tmp_path / "o2", seen_hashes=seen)
    assert len(first) == 1  # tiny image skipped
    assert second == []
    assert Image.open(first[0]).size == (300, 200)


def test_reference_citation_does_not_suppress_corpus_warning(tmp_path: Path):
    from docx import Document

    from pre_peer_checker.pipeline.orchestrator import run_verification
    from pre_peer_checker.warnings import WarningTag

    ms, corp = tmp_path / "ms", tmp_path / "corp"
    ms.mkdir()
    corp.mkdir()
    doc = Document()
    doc.add_paragraph("Figure 1. Clone size in eye discs (flies from Smith et al., 2012). n=12 (F).")
    doc.save(ms / "main.docx")
    photo = _photo(4, (256, 256))
    photo.save(ms / "panel.png")
    photo.save(corp / "old.png")

    result = run_verification([ms], corpus=[corp])
    assert result.artifacts.get("legend_citation_mentioned") is True
    assert result.artifacts.get("legend_reuse_statement") is False
    assert any(
        w.tag == WarningTag.IMAGE_REUSE and w.metadata.get("corpus_match") for w in result.warnings
    )


def test_corpus_panel_match_finds_degraded_reuse_in_screenshot(tmp_path: Path):
    pytest.importorskip("lightglue")
    from pre_peer_checker.imaging.corpus_scan import scan_against_corpus
    from pre_peer_checker.warnings import WarningTag

    reused, other_ms, unrelated = _photo(11), _photo(12), _photo(13)
    ms_dir, corpus_dir = tmp_path / "ms", tmp_path / "corpus"
    ms_dir.mkdir()
    corpus_dir.mkdir()
    _page([other_ms, reused, _photo(14)]).save(ms_dir / "supp_page.png")

    # Past-paper screenshot: the same photo downscaled and re-compressed, next to an
    # unrelated photo
    small = reused.resize((220, 220), Image.Resampling.BILINEAR)
    jpg = io.BytesIO()
    small.save(jpg, format="JPEG", quality=70)
    small = Image.open(io.BytesIO(jpg.getvalue())).convert("RGB")
    shot = Image.new("RGB", (560, 300), "white")
    shot.paste(unrelated.resize((220, 220)), (30, 40))
    shot.paste(small, (290, 40))
    shot.save(corpus_dir / "pubpeer_page.png")

    result = scan_against_corpus(
        [ms_dir / "supp_page.png"],
        [corpus_dir],
        legend_has_citation=False,
        prefer_dino=False,
        enable_partial=False,
    )
    panel_warns = [
        w for w in result.warnings if w.tag == WarningTag.IMAGE_REUSE and w.metadata.get("panel")
    ]
    assert len(panel_warns) == 1
    meta = panel_warns[0].metadata
    left, _top, right, _bottom = meta["panel_b"]
    assert 280 <= left <= 300 and right <= 520  # the degraded copy, not the unrelated photo
    assert all(abs(m["panel_b"][0] - left) < 40 for m in meta["panel_matches"])

    cited = scan_against_corpus(
        [ms_dir / "supp_page.png"],
        [corpus_dir],
        legend_has_citation=True,
        prefer_dino=False,
        enable_partial=False,
    )
    assert not [w for w in cited.warnings if w.metadata.get("panel")]
    assert any(m.get("matches") for m in cited.artifacts.get("cited_matches") or [])
