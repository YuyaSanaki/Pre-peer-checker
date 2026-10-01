"""Read spreadsheets and delimited text into typed cell grids.

One reader for every table format so the block / statistics layers see the same
cells regardless of how the lab saved the file:

* Excel ``.xlsx/.xlsm/.xls/.xlsb`` and ``.ods`` via ``python-calamine`` (fast; the
  format is sniffed from the bytes, so an ``.xls`` that is really xlsx, or the
  reverse, still opens). ``openpyxl`` / ``xlrd`` are fallbacks.
* ``.xls`` files that are really tab-separated text (instrument exports) and
  ``.csv/.tsv/.txt`` (optionally ``.gz``): encoding and delimiter are detected.

A grid row is a list of ``(kind, value)`` with kind ``"num" | "str" | "empty"``.
Coordinates start at A1, so a cell's position matches what the user sees in Excel.
"""

from __future__ import annotations

import csv
import datetime as _dt
import gzip
import io
import math
import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any

Cell = tuple[str, Any]

EXCEL_SUFFIXES = frozenset({".xlsx", ".xlsm", ".xls", ".xlsb", ".ods"})
TEXT_SUFFIXES = frozenset({".csv", ".tsv", ".txt"})

# Rows read per sheet. Plotted / per-sample tables are far smaller; huge sheets are
# omics matrices or per-cell dumps where only the first rows are needed to see the
# layout. ``Grid.total_rows`` keeps the true height.
DEFAULT_MAX_ROWS = 20000
_MAX_TEXT_BYTES = 64 * 1024 * 1024

_NUMERIC_STR_RE = re.compile(r"^\s*[-+]?(\d+(\.\d*)?|\.\d+)([eE][-+]?\d+)?\s*%?\s*$")
_THOUSANDS_RE = re.compile(r"^\s*[-+]?\d{1,3}(,\d{3})+(\.\d+)?\s*$")
_DECIMAL_COMMA_RE = re.compile(r"^\s*[-+]?\d+,\d+\s*$")
_MISSING_TOKENS = frozenset({"", "na", "n/a", "nan", "null", "none", "-", "--"})
_EXCEL_ERRORS = frozenset(
    {"#n/a", "#div/0!", "#value!", "#ref!", "#name?", "#num!", "#null!", "#getting_data"}
)


@dataclass
class Grid:
    path: Path
    sheet: str
    rows: list[list[Cell]]
    total_rows: int
    truncated: bool = False

    @property
    def width(self) -> int:
        return max((len(r) for r in self.rows), default=0)

    def cell(self, r: int, c: int) -> Cell:
        if r < 0 or r >= len(self.rows):
            return ("empty", None)
        row = self.rows[r]
        if c < 0 or c >= len(row):
            return ("empty", None)
        return row[c]


def parse_number(s: str, *, decimal_comma: bool = False) -> float | None:
    """Float for a numeric-looking string (``1.5``, ``1e-3``, ``12%``, ``1,234``)."""
    t = s.strip()
    if not t:
        return None
    if decimal_comma and _DECIMAL_COMMA_RE.match(t):
        t = t.replace(",", ".")
    elif _THOUSANDS_RE.match(t):
        t = t.replace(",", "")
    if not _NUMERIC_STR_RE.match(t):
        return None
    t = t.rstrip().rstrip("%")
    try:
        v = float(t)
    except ValueError:
        return None
    if math.isnan(v) or math.isinf(v):
        return None
    return v


def cell_kind(v: Any, *, decimal_comma: bool = False) -> Cell:
    if v is None:
        return ("empty", None)
    if isinstance(v, bool):
        return ("str", str(v))
    if isinstance(v, (int, float)):
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return ("empty", None)
        return ("num", float(v))
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time, _dt.timedelta)):
        return ("str", str(v))
    s = str(v).strip()
    if not s:
        return ("empty", None)
    num = parse_number(s, decimal_comma=decimal_comma)
    if num is not None:
        return ("num", num)
    if s.lower() in _MISSING_TOKENS or s.lower() in _EXCEL_ERRORS:
        return ("empty", None)
    return ("str", s)


def _trim_row(row: list[Cell]) -> list[Cell]:
    end = len(row)
    while end and row[end - 1][0] == "empty":
        end -= 1
    return row[:end]


def _trim_grid(rows: list[list[Cell]]) -> list[list[Cell]]:
    rows = [_trim_row(r) for r in rows]
    end = len(rows)
    while end and not rows[end - 1]:
        end -= 1
    return rows[:end]


def sniff_format(path: Path) -> str:
    """``"zip"`` (xlsx/ods), ``"ole"`` (legacy xls), ``"gzip"`` or ``"text"``."""
    try:
        with open(path, "rb") as fh:
            head = fh.read(8)
    except OSError:
        return "text"
    if head.startswith(b"PK\x03\x04"):
        return "zip"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "ole"
    if head.startswith(b"\x1f\x8b"):
        return "gzip"
    return "text"


def _excel_grids_calamine(path: Path, max_rows: int | None) -> list[Grid]:
    from python_calamine import CalamineWorkbook

    grids: list[Grid] = []
    with open(path, "rb") as fh:
        wb = CalamineWorkbook.from_filelike(fh)
        for name in wb.sheet_names:
            ws = wb.get_sheet_by_name(name)
            total = int(getattr(ws, "total_height", 0) or ws.height or 0)
            raw = ws.to_python(skip_empty_area=False, nrows=max_rows) if max_rows else (
                ws.to_python(skip_empty_area=False)
            )
            rows = _trim_grid([[cell_kind(v) for v in r] for r in raw])
            grids.append(
                Grid(path, str(name), rows, total_rows=max(total, len(rows)),
                     truncated=bool(max_rows and total > max_rows))
            )
    return grids


def _excel_grids_openpyxl(path: Path, max_rows: int | None) -> list[Grid]:
    import warnings

    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(path, read_only=True, data_only=True)
    grids: list[Grid] = []
    try:
        for ws in wb.worksheets:
            rows: list[list[Cell]] = []
            for i, row in enumerate(ws.iter_rows(min_row=1, min_col=1, values_only=True)):
                if max_rows and i >= max_rows:
                    break
                rows.append([cell_kind(v) for v in row])
            total = int(ws.max_row or len(rows))
            rows = _trim_grid(rows)
            grids.append(Grid(path, ws.title, rows, total_rows=max(total, len(rows)),
                              truncated=bool(max_rows and total > max_rows)))
    finally:
        wb.close()
    return grids


def _excel_grids_pandas(path: Path, max_rows: int | None) -> list[Grid]:
    import pandas as pd

    book = pd.read_excel(path, sheet_name=None, header=None, nrows=max_rows)
    grids: list[Grid] = []
    for name, df in book.items():
        rows = _trim_grid([[cell_kind(v) for v in r] for r in df.itertuples(index=False)])
        grids.append(Grid(path, str(name), rows, total_rows=len(rows)))
    return grids


def _read_excel_grids(path: Path, max_rows: int | None) -> list[Grid]:
    fmt = sniff_format(path)
    if fmt == "text":
        return _read_text_grids(path, max_rows)
    errors: list[str] = []
    readers = [_excel_grids_calamine]
    if fmt == "zip":
        readers.append(_excel_grids_openpyxl)
    readers.append(_excel_grids_pandas)
    for reader in readers:
        try:
            return reader(path, max_rows)
        except ImportError as exc:
            errors.append(f"{reader.__name__}: {exc}")
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{reader.__name__}: {type(exc).__name__}: {exc}")
    raise ValueError(f"Excel を読めません: {path.name}: " + " / ".join(errors))


def decode_text(data: bytes) -> str:
    """Decode text exports: UTF-8 (BOM), UTF-16, then a detected or legacy codepage."""
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    try:
        from charset_normalizer import from_bytes

        best = from_bytes(data[:200_000]).best()
        if best is not None and best.encoding:
            return data.decode(best.encoding, errors="replace")
    except Exception:  # noqa: BLE001
        pass
    for enc in ("cp932", "cp1252", "mac_roman"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def sniff_delimiter(lines: list[str]) -> str:
    """Delimiter whose per-line count is most consistent (and non-zero) on sample lines."""
    sample = [ln for ln in lines if ln.strip()][:200]
    if not sample:
        return ","
    best, best_score = ",", -1.0
    for d in ("\t", ",", ";", "|"):
        counts = [ln.count(d) for ln in sample]
        hits = [c for c in counts if c > 0]
        if not hits:
            continue
        mode = max(set(hits), key=hits.count)
        score = hits.count(mode) / len(sample) + 0.001 * mode
        if d == "\t":
            score += 0.05  # a tab inside a cell is rarer than a comma
        if score > best_score:
            best, best_score = d, score
    if best_score < 0 and any(re.search(r"\S\s{2,}\S", ln) for ln in sample):
        return "whitespace"
    return best


def _read_text_bytes(path: Path) -> bytes:
    opener = gzip.open if (path.suffix.lower() == ".gz" or sniff_format(path) == "gzip") else open
    with opener(path, "rb") as fh:
        return fh.read(_MAX_TEXT_BYTES)


def _read_text_grids(path: Path, max_rows: int | None) -> list[Grid]:
    text = decode_text(_read_text_bytes(path))
    lines = text.splitlines()
    delim = sniff_delimiter(lines[:400])
    decimal_comma = delim == ";"
    if delim == "whitespace":
        records = (re.split(r"\s{2,}|\t", ln.strip()) for ln in lines)
    else:
        records = csv.reader(io.StringIO(text), delimiter=delim)
    rows: list[list[Cell]] = []
    total = 0
    for rec in records:
        total += 1
        if max_rows and len(rows) >= max_rows:
            continue
        rows.append([cell_kind(v, decimal_comma=decimal_comma) for v in rec])
    rows = _trim_grid(rows)
    name = path.name
    for suf in (".gz",):
        if name.lower().endswith(suf):
            name = name[: -len(suf)]
    return [Grid(path, Path(name).stem, rows, total_rows=total,
                 truncated=bool(max_rows and total > max_rows))]


def _table_suffix(path: Path) -> str:
    name = path.name.lower()
    if name.endswith(".gz"):
        name = name[:-3]
    return Path(name).suffix


def read_grids_uncached(path: Path, *, max_rows: int | None = DEFAULT_MAX_ROWS) -> list[Grid]:
    suf = _table_suffix(path)
    if suf in EXCEL_SUFFIXES:
        return _read_excel_grids(path, max_rows)
    if suf in TEXT_SUFFIXES:
        return _read_text_grids(path, max_rows)
    raise ValueError(f"Unsupported table format: {path.suffix}")


@lru_cache(maxsize=128)
def _read_cached(path_str: str, mtime_ns: int, size: int, max_rows: int | None) -> tuple[Grid, ...]:
    return tuple(read_grids_uncached(Path(path_str), max_rows=max_rows))


def read_grids(path: Path | str, *, max_rows: int | None = DEFAULT_MAX_ROWS) -> list[Grid]:
    """Typed cell grids, one per sheet (text files give a single grid)."""
    p = Path(path)
    st = p.stat()
    return list(_read_cached(str(p.resolve()), st.st_mtime_ns, st.st_size, max_rows))


def _row_counts(row: list[Cell]) -> tuple[int, int]:
    n_num = sum(1 for k, _ in row if k == "num")
    n_str = sum(1 for k, _ in row if k == "str")
    return n_num, n_str


def find_header_row(rows: list[list[Cell]], *, search: int = 60) -> int | None:
    """Header of the first data table: a text row of ≥2 cells over ≥2 numeric rows.

    Skips preambles (instrument metadata, ``Detailed`` / ``=====`` banners, titles).
    Key/value metadata (``Chemistry | SYBR``) is not a header because the next
    rows are not numeric.
    """
    for i in range(min(search, len(rows))):
        n_num, n_str = _row_counts(rows[i])
        if n_str < 2 or n_num > n_str:
            continue
        width = len(rows[i])
        body = rows[i + 1 : i + 4]
        if len(body) < 2:
            continue
        ok = 0
        for r in body:
            b_num, b_str = _row_counts(r)
            if b_num >= 1 and b_num + b_str >= 2 and len(r) <= width + 1:
                ok += 1
        if ok >= min(2, len(body)):
            return i
    return None


def grid_to_frame(grid: Grid, header_row: int | None = 0):
    """DataFrame from a grid; ``header_row=None`` keeps every row as data."""
    import pandas as pd

    rows = grid.rows
    width = grid.width
    values = [[v if k != "empty" else None for k, v in r] + [None] * (width - len(r)) for r in rows]
    if header_row is None or header_row >= len(values):
        return pd.DataFrame(values)
    header = values[header_row]
    names: list[str] = []
    seen: dict[str, int] = {}
    for j, h in enumerate(header):
        base = f"col_{j}" if h is None or str(h).strip() == "" else str(h).strip()
        if isinstance(h, float) and h.is_integer():
            base = str(int(h))
        k = seen.get(base, 0)
        seen[base] = k + 1
        names.append(base if k == 0 else f"{base}.{k}")
    body = values[header_row + 1 :]
    df = pd.DataFrame(body, columns=names)
    df = df.dropna(how="all")
    df = df.loc[:, [c for c in df.columns if not (c.startswith("col_") and df[c].isna().all())]]
    for c in df.columns:
        conv = pd.to_numeric(df[c], errors="coerce")
        if conv.notna().sum() == df[c].notna().sum():
            if len(conv) and conv.notna().all() and (conv % 1 == 0).all() and (conv.abs() < 2**53).all():
                conv = conv.astype("int64")
            df[c] = conv
    return df.reset_index(drop=True)


def is_table_path(path: Path | str) -> bool:
    suf = _table_suffix(Path(path))
    return suf in EXCEL_SUFFIXES or suf in {".csv", ".tsv"}
