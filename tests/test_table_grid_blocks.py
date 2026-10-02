"""Grid reader + generalized Source Data blocks on synthetic workbooks / text tables."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook

from pre_peer_checker.data.figure_refs import (
    expand_panels,
    parse_figure_label,
    sheet_figure_label,
    valid_panel_run,
    workbook_uses_short_sheet_names,
)
from pre_peer_checker.data.source_data_blocks import (
    parse_source_data_blocks,
    parse_source_data_bundle,
)
from pre_peer_checker.data.table_grid import find_header_row, read_grids, sniff_delimiter
from pre_peer_checker.engine.source_data_checks import (
    match_source_block_for_panel,
    warnings_from_source_data_panels,
)
from pre_peer_checker.parsers.legend_struct import PanelN


def _save(path: Path, sheets: dict[str, list[list]]) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)
    return path


# --- figure labels -----------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "number", "panels", "extended"),
    [
        ("Fig. 3e", "3", ("e",), False),
        ("Figure 2c-e", "2", ("c", "d", "e"), False),
        ("Extended Data Fig. 5a", "5", ("a",), True),
        ("Source data for Extended Fig 1", "1", (), True),
        ("Supplementary Figure S4b", "S4", ("b",), True),
        ("Fid.1e", "1", ("e",), False),
    ],
)
def test_parse_figure_label(text, number, panels, extended):
    lbl = parse_figure_label(text)
    assert lbl is not None
    assert (lbl.number, lbl.panels, lbl.extended) == (number, panels, extended)


@pytest.mark.parametrize("text", ["Fig 1 Lifespan", "data", "qPCR results", "graph"])
def test_words_are_not_panel_runs(text):
    lbl = parse_figure_label(text)
    assert lbl is None or lbl.panels == ()


def test_panel_runs():
    assert valid_panel_run("bf") and valid_panel_run("acdg")
    assert not valid_panel_run("data")
    assert expand_panels("c–e") == ("c", "d", "e")


def test_short_sheet_names():
    assert workbook_uses_short_sheet_names(["2a", "2bf", "2c"])
    assert not workbook_uses_short_sheet_names(["Sheet1", "raw", "2a"])
    # lab numbering: dose series and zero-padded plate numbers are not panels
    assert not workbook_uses_short_sheet_names(["Sheet1", "1x", "2x", "3x", "4x"])
    assert not workbook_uses_short_sheet_names(["sum", "1P", "09P", "08P", "07P"])
    assert workbook_uses_short_sheet_names(["1c", "2c", "3c"], series_guard=False)
    lbl = sheet_figure_label("3bf", file_label=None, short_names_ok=True)
    assert lbl is not None and lbl.number == "3" and lbl.panels == ("b", "f")
    assert sheet_figure_label("3bf", file_label=None, short_names_ok=False) is None


# --- grid reader -------------------------------------------------------------


def test_xls_suffix_with_xlsx_content(tmp_path):
    src = _save(tmp_path / "a.xlsx", {"S": [["Ctrl", "KO"], [1, 2], [3, 4]]})
    dst = tmp_path / "renamed.xls"
    shutil.copy(src, dst)
    grids = read_grids(dst)
    assert grids and grids[0].rows[1][0] == ("num", 1.0)


def test_text_table_decimal_comma_and_preamble(tmp_path):
    p = tmp_path / "export.csv"
    p.write_text(
        "Instrument: X\nRun date: 2020-01-01\n\nSample;Value;Group\nA1;1,5;ctrl\nA2;2,25;ctrl\nA3;3,0;ko\n",
        encoding="utf-8",
    )
    grid = read_grids(p)[0]
    hdr = find_header_row(grid.rows)
    assert hdr is not None and grid.rows[hdr][0] == ("str", "Sample")
    assert grid.rows[hdr + 2][1] == ("num", 2.25)


def test_imaris_style_preamble(tmp_path):
    p = tmp_path / "Volume.csv"
    p.write_text(
        "Volume\n====================\n\nVolume,Unit,Category,ID\n"
        "10.5,um^3,Surface,1\n11.2,um^3,Surface,2\n9.8,um^3,Surface,3\n",
        encoding="utf-8",
    )
    grid = read_grids(p)[0]
    hdr = find_header_row(grid.rows)
    assert hdr is not None and grid.rows[hdr][0] == ("str", "Volume")
    assert grid.rows[hdr + 1][0] == ("num", 10.5)


def test_sniff_delimiter():
    assert sniff_delimiter(["a\tb\tc", "1\t2\t3"]) == "\t"
    assert sniff_delimiter(["a;b;c", "1,5;2;3"]) == ";"


# --- Source Data blocks ------------------------------------------------------


def test_sheet_per_panel_workbook(tmp_path):
    p = _save(
        tmp_path / "Source_Data_Fig2.xlsx",
        {
            "2a": [["Ctrl", "KO"], [1, 2], [3, 4], [5, 6]],
            "2bf": [["Ctrl", "KO"], [1, 2], [3, 4]],
        },
    )
    blocks = parse_source_data_blocks(p)
    by_sheet = {b.sheet: b for b in blocks}
    assert by_sheet["2a"].figure.number == "2"
    assert by_sheet["2a"].layout == "wide"
    assert [len(v) for _, v in by_sheet["2a"].groups] == [3, 3]
    assert by_sheet["2bf"].all_panels == ("b", "f")


def test_long_format_groups(tmp_path):
    rows = [["Fig. 1c"], ["Genotype", "Value"]]
    rows += [["WT", v] for v in (1.0, 2.0, 3.0)]
    rows += [["KO", v] for v in (4.0, 5.0)]
    blocks = parse_source_data_blocks(_save(tmp_path / "sd.xlsx", {"Sheet1": rows}))
    assert len(blocks) == 1
    b = blocks[0]
    assert b.layout == "long"
    assert dict((g, len(v)) for g, v in b.groups) == {"WT": 3, "KO": 2}
    assert b.legend_n_conflict(3, "WT") is None
    assert b.legend_n_conflict(3) == 2


def test_time_course_is_not_sample_n(tmp_path):
    rows = [["Fig. 2b"], ["Days", "Ctrl", "MetR"]]
    rows += [[d, 100 - d, 100 - d / 2] for d in range(0, 50, 5)]
    b = parse_source_data_blocks(_save(tmp_path / "sd.xlsx", {"S": rows}))[0]
    assert b.layout == "xy" and not b.n_comparable
    assert b.legend_n_conflict(6) is None


def _km_percent(n: int, events: list[tuple[int, int]]) -> list[float]:
    """Percent survival after each (deaths, censored-after) step, 6 decimals like Prism."""
    at_risk, s, out = n, 1.0, [100.0]
    for deaths, censored in events:
        if deaths:
            s *= 1 - deaths / at_risk
        at_risk -= deaths + censored
        out.append(round(s * 100, 6))
    return out


def test_km_cohort_recovers_n_with_censoring():
    from pre_peer_checker.data.curve_n import km_cohort_size

    events = [(0, 0), (2, 0), (1, 1), (3, 0), (0, 2), (5, 0), (4, 1), (8, 0), (6, 0), (4, 0)]
    assert km_cohort_size(_km_percent(37, events)) == 37
    # censoring before the first death leaves the curve unchanged: n is a minimum
    assert km_cohort_size(_km_percent(40, [(0, 3), *events])) == 37


def test_survival_curve_gives_analysis_n(tmp_path):
    ctrl = _km_percent(37, [(0, 0), (2, 0), (1, 1), (3, 0), (5, 0), (8, 2), (10, 0), (5, 0)])
    metr = _km_percent(41, [(0, 0), (1, 0), (2, 0), (4, 1), (6, 0), (9, 0), (10, 2), (6, 0)])
    rows = [["Fig. 2b"], ["Lifespan"], ["days", "Ctrl", "MetR"]]
    rows += [[d * 7, c, m] for d, (c, m) in enumerate(zip(ctrl, metr))]
    b = parse_source_data_blocks(_save(tmp_path / "sd.xlsx", {"Fig.2b": rows}))[0]
    assert b.layout == "xy" and b.n_comparable
    assert b.group_ns == [("Ctrl", 37, False), ("MetR", 41, False)]
    assert b.legend_n_conflict(30, "Ctrl") == 37
    assert b.legend_n_conflict(39, "Ctrl") is None
    pn = PanelN(figure="Fig. 2", panel="b", n=30, context="n = 30 flies", group="Ctrl")
    ws = warnings_from_source_data_panels([b], [pn])
    assert len(ws) == 1 and ws[0].metadata["n_authority"] == "survival_curve_min_cohort"


def test_alive_counts_give_exact_n(tmp_path):
    rows = [["Fig. 3a"], ["Survival"], ["day", "WT", "KO"]]
    rows += [[d, a, b] for d, (a, b) in enumerate([(20, 18), (20, 15), (17, 11), (12, 4), (5, 0)])]
    b = parse_source_data_blocks(_save(tmp_path / "sd.xlsx", {"S": rows}))[0]
    assert b.group_ns == [("WT", 20, True), ("KO", 18, True)]
    assert b.legend_n_conflict(19, "WT") == 20


def test_numeric_header_label_is_a_column(tmp_path):
    rows = [["Fig. 2j"], ["Lifespan"], ["days", 100, "Early 10%", "Lifelong 10%"]]
    rows += [[d, 100 - d, 100 - 2 * d, 100 - 3 * d] for d in range(0, 30, 5)]
    b = parse_source_data_blocks(_save(tmp_path / "sd.xlsx", {"S": rows}))[0]
    assert [c.header for c in b.columns] == ["days", "100", "Early 10%", "Lifelong 10%"]


def test_replicate_rows_pool_into_stacked_header_groups(tmp_path):
    rows = [
        ["Fig. 4d"],
        ["Primed", "Naive"],
        ["shX", "shX"],
        ["Experiment 1", "Experiment 1"],
        [10, 20],
        [11, 21],
        [12, 22],
        ["Experiment 2", "Experiment 2"],
        [13, 23],
        [14, 24],
    ]
    blocks = parse_source_data_blocks(_save(tmp_path / "sd.xlsx", {"S": rows}))
    assert len(blocks) == 1
    assert [(g, len(v)) for g, v in blocks[0].groups] == [("Primed shX", 5), ("Naive shX", 5)]


def test_renamed_copy_takes_extended_label(tmp_path):
    sheets = {"3bf": [["Ctrl", "KO"], [1, 2], [3, 4]], "3c": [["A", "B"], [5, 6], [7, 8]]}
    _save(tmp_path / "Source_Data_ED_Fig3.xlsx", sheets)
    _save(tmp_path / "MOESM11_ESM.xlsx", sheets)
    blocks = parse_source_data_bundle(sorted(tmp_path.glob("*.xlsx")))
    assert blocks and all(b.figure.extended for b in blocks)


def test_sub_tables_with_same_n_link(tmp_path):
    rows = [["Fig. 4i"], ["miR-1"], ["A", "B"], [1, 2], [3, 4], [], ["miR-2"], ["A", "B"], [5, 6], [7, 8]]
    blocks = parse_source_data_blocks(_save(tmp_path / "sd.xlsx", {"S": rows}))
    assert len(blocks) == 2
    pn = PanelN(panel="I", n=2, figure="Figure 4", context="n = 2")
    block, covered = match_source_block_for_panel(pn, blocks)
    assert covered and block is not None


def test_textclipping_plist_decoding(tmp_path):
    import plistlib

    from pre_peer_checker.parsers.r_residue import find_residue_files, parse_textclipping

    def clip(name: str, text: str) -> Path:
        p = tmp_path / name
        p.write_bytes(
            plistlib.dumps(
                {"UTI-Data": {"public.utf8-plain-text": text}}, fmt=plistlib.FMT_BINARY
            )
        )
        return p

    welch = clip(
        "a.textclipping",
        "data:  x and y\nt = -2.82, df = 8.9583, p-value = 0.02013\n"
        "alternative hypothesis: true difference in means is not equal to 0\n"
        "sample estimates:\nmean of x mean of y \n 74.77 81.00\n",
    )
    dwass = clip("b.textClipping", "Group sizes: 9 10 7 11 \nUsing the Asymptotic method:\nDwass\n")
    st = parse_textclipping(welch)
    assert st.kind == "welch_t" and st.df == pytest.approx(8.9583)
    assert st.means == (74.77, 81.0) and st.p_value == pytest.approx(0.02013)
    assert parse_textclipping(dwass).group_sizes == (9, 10, 7, 11)
    assert {p.name for p in find_residue_files(tmp_path)} == {welch.name, dwass.name}


def test_legend_or_range_is_not_exact():
    from pre_peer_checker.parsers.legend_struct import _N_MENTION_RE

    m = _N_MENTION_RE.search("n = 2 or 3 independent experiments")
    assert m and (m.group(1), m.group(2)) == ("2", "3")


def test_legend_n_warning_per_group(tmp_path):
    rows = [["Fig. 1c"], ["Genotype", "Value"]]
    rows += [["WT", v] for v in (1.0, 2.0, 3.0)]
    rows += [["KO", v] for v in (4.0, 5.0)]
    blocks = parse_source_data_blocks(_save(tmp_path / "sd.xlsx", {"S": rows}))
    pn = PanelN(panel="C", n=3, figure="Figure 1", context="n = 3 per genotype")
    warns = warnings_from_source_data_panels(blocks, [pn])
    assert len(warns) == 1 and warns[0].metadata["data_n"] == 2


def test_count_table_sums_counts_per_group(tmp_path):
    from pre_peer_checker.data.group_vectors import extract_group_vectors

    header = ["Genotype", "Class", "label", "number", "percent"]
    path = _save(tmp_path / "phenotype.xlsx", {
        "eye": [header, ["cont", "I", 1, 53, 100], ["cont", "II", 2, 0, 0],
                ["mut", "I", 3, 10, 18], ["mut", "II", 4, 45, 82]],
        "wing": [header, ["cont", "I", 1, 130, 93], ["cont", "II", 2, 10, 7],
                 ["mut", "I", 3, 34, 39], ["mut", "II", 4, 54, 61]],
    })
    got = {(v.sheet, v.group_key, v.n) for v in extract_group_vectors(path)}
    assert got == {("eye", "cont", 53), ("eye", "mut", 55), ("wing", "cont", 140), ("wing", "mut", 88)}


def test_tidy_group_sheet_wins_over_ratio_sheet(tmp_path):
    from pre_peer_checker.data.group_vectors import extract_group_vectors

    tidy = [["id", "clone_vol", "label", "genotype", "date"]]
    tidy += [[i, 5.0 + i, 1, "ctrl", 20200101] for i in range(4)]
    tidy += [[i, 9.0 + i, 2, "mut", 20200101] for i in range(4, 10)]
    ratio = [[None] * 4, ["disc ID", "pixel count", "clone", "ratio (%)"]]
    ratio += [[i, 1000.0, 50.0 + i, 5.0 + i] for i in range(4)]
    path = _save(tmp_path / "CloneRatio.xlsx", {"Sheet1": tidy, "ctrl": ratio})
    got = {(v.group_key, v.n) for v in extract_group_vectors(path)}
    assert got == {("ctrl", 4), ("mut", 6)}
