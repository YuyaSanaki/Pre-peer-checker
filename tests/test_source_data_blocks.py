"""Journal Source Data workbooks: titled side-by-side blocks per figure panel."""

from __future__ import annotations

import math
from pathlib import Path

from openpyxl import Workbook

from pre_peer_checker.data.group_vectors import extract_group_vectors
from pre_peer_checker.data.source_data_blocks import parse_source_data_blocks
from pre_peer_checker.data.stats_recalc import analyze_table_file
from pre_peer_checker.engine.n_and_names import is_plot_quant_table
from pre_peer_checker.engine.source_data_checks import (
    warnings_from_source_data_panels,
    warnings_from_source_data_reuse,
    warnings_from_source_data_summaries,
)
from pre_peer_checker.parsers.legend_struct import PanelN

SCRIB = [14.08, 13.32, 11.16, 10.06, 11.73, 11.16, 11.43, 12.33, 13.08, 12.32]
RNAI = [69.46, 56.39, 56.50, 66.90, 68.37, 71.56, 77.21, 75.26, 75.64, 63.54, 45.52]
EDU = [0.5, 1.8, 2.4, 1.7, 0.2, 2.7, 0.9, 0.7, 0.9, 1.7, 1.0, 2.8]


def _mean(v):
    return sum(v) / len(v)


def _sd(v, ddof):
    m = _mean(v)
    return math.sqrt(sum((x - m) ** 2 for x in v) / (len(v) - ddof))


def _block(ws, row, col, title, headers, columns, *, summary=None, fig=None):
    """Write one block; ``summary`` = {"Mean": x, "SD": y, "SEM": z} under the last column."""
    if fig:
        ws.cell(row=row - 2, column=col, value=fig)
    ws.cell(row=row - 1, column=col, value=title)
    ws.cell(row=row, column=col, value="n")
    for j, h in enumerate(headers, start=1):
        ws.cell(row=row, column=col + j, value=h)
    n = len(columns[0])
    for i in range(n):
        ws.cell(row=row + 1 + i, column=col, value=i + 1)
        for j, colvals in enumerate(columns, start=1):
            ws.cell(row=row + 1 + i, column=col + j, value=colvals[i])
    for k, (label, val) in enumerate((summary or {}).items()):
        ws.cell(row=row + 1 + n + k, column=col + len(columns) - 1, value=label)
        ws.cell(row=row + 1 + n + k, column=col + len(columns), value=val)


def _fig3(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    _block(
        ws, 5, 2, "scrib vs wild-type (Fig. 3b)", ["Area", "positive %"],
        [[v * 1000 for v in SCRIB], SCRIB],
        summary={"Mean": _mean(SCRIB), "SD": _sd(SCRIB, 0), "SEM": _sd(SCRIB, 1) / math.sqrt(10)},
        fig="Fig. 3e",
    )
    _block(
        ws, 25, 2, "scrib+RNAi vs wild-type (Fig. 3c)", ["Area", "positive %"],
        [[v * 1000 for v in RNAI], RNAI],
        summary={"Mean": _mean(RNAI), "SD": _sd(RNAI, 1)},
    )
    wb.save(path)


def _fig4(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    _block(
        ws, 5, 2, "scrib vs wild-type (Fig. 4b)", ["Area", "Relative size (%)"],
        [[v * 1000 for v in SCRIB], SCRIB],
        summary={"Mean": _mean(SCRIB), "SD": _sd(SCRIB, 1)},
        fig="Fig. 4h",
    )
    edu_sub = EDU[1:]
    _block(
        ws, 5, 7, "wild-type vs wild-type (Fig. 4m)", ["EdU/area"],
        [EDU],
        summary={"Mean": _mean(edu_sub), "SD": _sd(edu_sub, 1)},
        fig="Fig. 4p",
    )
    _block(ws, 30, 7, "scrib vs wild-type (Fig. 4n)", ["EdU/area"], [EDU[:8]])
    _block(ws, 30, 2, "scrib+RNAi, EGFR-RNAi vs wild-type (Fig. 4d)", ["Rel (%)"], [RNAI[:9]])
    _block(ws, 50, 2, "scrib+RNAi, Wts vs wild-type (Fig. 4d)", ["Rel (%)"], [SCRIB[:7]])
    wb.save(path)


def _blocks(tmp_path: Path):
    a, b = tmp_path / "MOESM41_ESM.xlsx", tmp_path / "MOESM43_ESM.xlsx"
    _fig3(a)
    _fig4(b)
    return parse_source_data_blocks(a) + parse_source_data_blocks(b), a, b


def test_parses_side_by_side_blocks_with_inherited_figure(tmp_path: Path):
    blocks, _, _ = _blocks(tmp_path)
    by_title = {blk.title: blk for blk in blocks}
    assert len(blocks) == 7
    rnai3 = by_title["scrib+RNAi vs wild-type (Fig. 3c)"]
    assert rnai3.figure is not None and rnai3.figure.label() == "Fig. 3e"
    assert rnai3.has_index and rnai3.n == len(RNAI)
    assert rnai3.primary.header == "positive %"
    edu = by_title["wild-type vs wild-type (Fig. 4m)"]
    assert edu.figure.label() == "Fig. 4p" and edu.n == len(EDU)
    assert by_title["scrib vs wild-type (Fig. 4n)"].figure.label() == "Fig. 4p"


def test_source_data_feeds_group_vectors_and_plot_tables(tmp_path: Path):
    _, a, _ = _blocks(tmp_path)
    assert is_plot_quant_table(a)
    vecs = {v.group_key: v for v in extract_group_vectors(a)}
    assert vecs["Fig. 3e | scrib vs wild-type (Fig. 3b)"].n == len(SCRIB)
    assert vecs["Fig. 3e | scrib+RNAi vs wild-type (Fig. 3c)"].n == len(RNAI)
    stats = analyze_table_file(a)
    assert stats.anova is None and not stats.pairwise


def test_plain_excel_is_not_source_data(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["label", "value"])
    for i in range(6):
        ws.append(["wt", float(i)])
    p = tmp_path / "plain.xlsx"
    wb.save(p)
    assert parse_source_data_blocks(p) == []
    assert not is_plot_quant_table(p)


def test_cross_figure_reuse(tmp_path: Path):
    blocks, _, _ = _blocks(tmp_path)
    warns = warnings_from_source_data_reuse(blocks)
    assert len(warns) == 1
    w = warns[0]
    assert w.metadata["pattern_id"] == "P-SOURCE-DATA-CROSS-FIGURE-REUSE"
    assert w.metadata["same_condition"] is True
    assert "Fig. 3e" in w.title and "Fig. 4h" in w.title
    assert warnings_from_source_data_reuse(blocks, disclosure="explicit") == []


def test_summary_rows_population_sd_and_excluded_row(tmp_path: Path):
    blocks, _, _ = _blocks(tmp_path)
    warns = warnings_from_source_data_summaries(blocks)
    pids = [w.metadata["pattern_id"] for w in warns]
    assert pids.count("P-SOURCE-DATA-POPULATION-SD") == 1
    pop = next(w for w in warns if w.metadata["pattern_id"] == "P-SOURCE-DATA-POPULATION-SD")
    assert "Fig. 3b" in pop.reason and "Fig. 4b" not in pop.reason
    sub = [w for w in warns if w.metadata["pattern_id"] == "P-SOURCE-DATA-SUMMARY-MISMATCH"]
    assert len(sub) == 1
    assert sub[0].metadata["excluded_rows"] == [1]
    assert "Fig. 4m" in sub[0].title


def test_legend_n_and_duplicate_panel_ref(tmp_path: Path):
    blocks, _, _ = _blocks(tmp_path)
    panel_ns = [
        PanelN(panel="P", n=12, figure="Figure 4", context="m (n = 12)", group="m"),
        PanelN(panel="P", n=30, figure="Figure 4", context="n (n = 30)", group="n"),
        PanelN(panel="H", n=10, figure="Figure 4", context="b (n = 10)", group="b"),
        PanelN(panel="H", n=9, figure="Figure 4", context="d (n = 9)", group="d"),
        PanelN(panel="H", n=7, figure="Figure 4", context="g (n = 7)", group="g"),
    ]
    warns = warnings_from_source_data_panels(blocks, panel_ns)
    n_warns = [w for w in warns if w.metadata["pattern_id"] == "P-SOURCE-DATA-LEGEND-N"]
    assert len(n_warns) == 1
    assert n_warns[0].metadata["legend_n"] == 30 and n_warns[0].metadata["data_n"] == 8
    dup = [w for w in warns if w.metadata["pattern_id"] == "P-SOURCE-DATA-PANEL-REF-DUP"]
    assert len(dup) == 1
    assert "(g) n=7" in dup[0].reason
