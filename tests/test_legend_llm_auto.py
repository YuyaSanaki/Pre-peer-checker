"""Legend LLM auto mode: rules first, the LLM reads only figures the rules left unread."""

from __future__ import annotations

import json
from pathlib import Path

from pre_peer_checker.llm.backend import CallableBackend
from pre_peer_checker.llm.legend_extract import (
    extract_legends_with_backend,
    legend_sample_sizes,
    normalize_legend_llm_mode,
    summarize_legend_llm_meta,
)


def test_normalize_mode_accepts_legacy_bool() -> None:
    assert normalize_legend_llm_mode(True) == "on"
    assert normalize_legend_llm_mode(False) == "off"
    assert normalize_legend_llm_mode(None) == "off"
    assert normalize_legend_llm_mode("auto") == "auto"
    assert normalize_legend_llm_mode("ON") == "on"
    assert normalize_legend_llm_mode("nonsense") == "off"


def test_legend_sample_sizes() -> None:
    assert legend_sample_sizes("n = 19 (a), 16 (d), and 13 (g)") == {19, 16, 13}
    assert legend_sample_sizes(
        "N = 15 (day 3), 93 (day 6) aggregates from 3 independent experiments."
    ) == {15, 93}
    assert legend_sample_sizes("Data are from three independent experiments.") == {3}
    assert legend_sample_sizes("Twelve images per condition; 12 mice per group.") == {12}
    assert legend_sample_sizes("(n =\u200b 5, P < 0.05)") == {5}
    assert legend_sample_sizes("Scale bars, 20 µm.") == set()


def _two_figure_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Figure 1. Clone size in control discs (A). n = 6 (A).")
    doc.add_paragraph("Figure 2. Body weight (A) was measured in 12 mice per group.")
    doc.save(path)
    return path


def _fake_llm(calls: list[str]):
    def fn(prompt: str) -> str:
        calls.append(prompt)
        if "[#1]" in prompt:  # auto: assignment of the tagged n
            return json.dumps({"tags": [{"id": 1, "n": 12, "panels": ["A"], "group": ""}]})
        return json.dumps(
            {"figure": "Figure 2", "panels": [{"panel": "A", "n": 12, "groups": [], "notes": ""}]}
        )

    return fn


def test_auto_calls_llm_only_for_unread_figures(tmp_path: Path) -> None:
    docx = _two_figure_docx(tmp_path / "m.docx")
    calls: list[str] = []
    legs, meta = extract_legends_with_backend(
        docx, mode="auto", backend=CallableBackend(_fake_llm(calls))
    )
    assert meta["legend_llm_mode"] == "auto"
    assert meta["n_llm_calls"] == 1 == len(calls)
    assert "Figure 2" in calls[0]
    by_fig = {leg.figure: leg for leg in legs}
    assert by_fig["Figure 1"].extractor == "rules"
    assert by_fig["Figure 2"].extractor != "rules"
    assert any(p.n == 12 for p in by_fig["Figure 2"].panels)

    status = summarize_legend_llm_meta([{"path": str(docx), **meta}])
    assert status["status"] == "active" and status["n_llm_calls"] == 1
    assert "1/2" in status["message"]


def test_auto_skips_llm_when_rules_read_everything(tmp_path: Path) -> None:
    from docx import Document

    docx = tmp_path / "m.docx"
    doc = Document()
    doc.add_paragraph("Figure 1. Clone size (A). n = 6 (A).")
    doc.save(docx)
    calls: list[str] = []
    _legs, meta = extract_legends_with_backend(
        docx, mode="auto", backend=CallableBackend(_fake_llm(calls))
    )
    assert calls == [] and meta["n_llm_calls"] == 0
    status = summarize_legend_llm_meta([{"path": str(docx), **meta}])
    assert status["status"] == "auto_skipped"


def test_on_mode_reads_every_figure(tmp_path: Path) -> None:
    docx = _two_figure_docx(tmp_path / "m.docx")
    calls: list[str] = []
    _legs, meta = extract_legends_with_backend(
        docx, mode="on", backend=CallableBackend(_fake_llm(calls))
    )
    assert meta["n_llm_calls"] == 2 == len(calls)
